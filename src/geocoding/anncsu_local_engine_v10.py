#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Aste Easy - ANNCSU benchmark v2

Scopo:
- NON modifica Supabase.
- NON chiama Google Maps.
- Usa lo storico google_geocode_log_v8.jsonl come ground truth di benchmark.
- Risolve i Comuni tramite l'elenco ufficiale ISTAT.
- Legge in streaming STRAD_ITA_*.csv e INDIR_ITA_*.csv.
- Tiene solo le strade/accessi utili ai record benchmark.
- Produce CSV + JSON di riepilogo per misurare quanto ANNCSU possa sostituire Google.

Richiede solo Python standard library. Se openpyxl è presente lo usa per l'xlsx ISTAT,
altrimenti include un lettore XLSX minimale basato su zip/xml.
"""

from __future__ import annotations

import argparse
import csv
import difflib
import json
import math
import os
import re
import sys
import time
import unicodedata
import urllib.request
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
import xml.etree.ElementTree as ET

VERSION = "10.0.2-local-engine"
ISTAT_XLSX_URL = "https://www.istat.it/storage/codici-unita-amministrative/Elenco-comuni-italiani.xlsx"

STOP_TOKENS = {
    "a", "ad", "al", "allo", "alla", "ai", "agli", "alle",
    "da", "dal", "dallo", "dalla", "dai", "dagli", "dalle",
    "di", "del", "dello", "della", "dei", "degli", "delle",
    "in", "nel", "nello", "nella", "nei", "negli", "nelle",
    "su", "sul", "sullo", "sulla", "sui", "sugli", "sulle",
    "e", "ed", "il", "lo", "la", "i", "gli", "le", "un", "uno", "una",
    "via", "viale", "v", "piazza", "pza", "pzza", "corso", "cso",
    "strada", "str", "vicolo", "vico", "largo", "salita", "traversa",
    "contrada", "cda", "localita", "loc", "frazione", "fraz", "borgata",
    "regione", "quartiere", "zona", "comunale", "provinciale", "statale",
    "numero", "num", "n", "civico", "civ", "snc", "s",
    "italia"
}


NUMBER_WORD_TO_CANON = {
    "uno":"1","due":"2","tre":"3","quattro":"4","cinque":"5","sei":"6","sette":"7","otto":"8","nove":"9","dieci":"10",
    "undici":"11","dodici":"12","tredici":"13","quattordici":"14","quindici":"15","sedici":"16","diciassette":"17",
    "diciotto":"18","diciannove":"19","venti":"20","ventuno":"21","ventidue":"22","ventitre":"23","ventitré":"23",
    "ventiquattro":"24","venticinque":"25","ventisei":"26","ventisette":"27","ventotto":"28","ventinove":"29",
    "trenta":"30","trentuno":"31",
    "primo":"1","prima":"1","secondo":"2","seconda":"2","terzo":"3","terza":"3","quarto":"4","quarta":"4",
    "quinto":"5","quinta":"5",
}
MONTH_TOKENS = {
    "gennaio","febbraio","marzo","aprile","maggio","giugno","luglio","agosto","settembre","ottobre","novembre","dicembre"
}

def roman_to_int_token(tok: str) -> Optional[int]:
    vals = {"i":1,"v":5,"x":10,"l":50,"c":100}
    t = (tok or "").lower()
    if not t or any(ch not in vals for ch in t):
        return None
    total = 0
    prev = 0
    for ch in reversed(t):
        v = vals[ch]
        if v < prev:
            total -= v
        else:
            total += v
            prev = v
    if not (1 <= total <= 100):
        return None
    return total

def canonical_street_tokens(value: object) -> List[str]:
    raw = norm_text(value).split()
    out: List[str] = []
    for i, t in enumerate(raw):
        if t in STOP_TOKENS:
            continue
        if t in NUMBER_WORD_TO_CANON:
            out.append(NUMBER_WORD_TO_CANON[t])
            continue
        if t.isdigit():
            out.append(str(int(t)))
            continue
        # Roman numerals of length >=2 are usually dates/names (XXV Aprile, XX Settembre, XXIII).
        # Single I/V/X/C are ambiguous with initials: convert only before a month or at the end.
        ri = roman_to_int_token(t)
        if ri is not None:
            nxt = raw[i + 1] if i + 1 < len(raw) else ""
            if len(t) >= 2 or nxt in MONTH_TOKENS or i == len(raw) - 1:
                out.append(str(ri))
                continue
        # Single-letter initials are deliberately ignored: G. Mazzini -> Mazzini.
        if len(t) < 2:
            continue
        out.append(t)
    return out

def alias_street_score(src_norm: str, odonimo: str) -> float:
    """Supplementary conservative matcher for aliases/initials/date numerals."""
    a = canonical_street_tokens(src_norm)
    b = canonical_street_tokens(odonimo)
    if not a or not b:
        return 0.0

    # Road codes are highly distinctive.
    def rc(v: str):
        n = norm_text(v).replace(" ", "")
        return set(re.findall(r"(?:sp|ss|sr|sc|a)\d{1,4}", n))
    ac, bc = rc(src_norm), rc(odonimo)
    if ac:
        if ac & bc:
            return 1.0
        # An explicit road code should not be "rescued" by an unrelated named street.
        if not bc and len(a) <= 1:
            return 0.0

    def sim(x: str, y: str) -> float:
        return 1.0 if x == y else difflib.SequenceMatcher(None, x, y).ratio()

    # If PVP effectively supplies only the distinctive surname/toponym, require it
    # to match the final official token. This avoids "S. Carlo" -> "Carlo Porta".
    if len(a) == 1:
        s = sim(a[0], b[-1])
        if s >= 0.92:
            return 0.96
        if s >= 0.84:
            return 0.84
        return s * 0.55

    bests = [max(sim(x, y) for y in b) for x in a]
    matched = sum(v >= 0.84 for v in bests)
    coverage = matched / len(a)
    tail = sim(a[-1], b[-1])

    if coverage == 1.0:
        return 0.98 if tail >= 0.84 else 0.90
    if coverage >= (2 / 3) and tail >= 0.84:
        return 0.90
    if coverage >= 0.50 and matched >= 2:
        return 0.82
    seq = difflib.SequenceMatcher(None, " ".join(a), " ".join(b)).ratio()
    return max(0.62 * coverage + 0.22 * tail, seq * 0.62)



def norm_text(value: object) -> str:
    s = "" if value is None else str(value)
    s = s.replace("’", "'").replace("`", "'")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_name(value: object) -> str:
    return norm_text(value)


def tokens(value: object, *, drop_stop: bool = True) -> List[str]:
    out = norm_text(value).split()
    if drop_stop:
        out = [t for t in out if t not in STOP_TOKENS and len(t) >= 2]
    return out


def normalize_civic(value: object) -> str:
    s = norm_text(value).replace(" ", "")
    return re.sub(r"[^a-z0-9]", "", s)


def safe_float(value: object) -> Optional[float]:
    if value is None:
        return None
    s = str(value).strip().replace(",", ".")
    if not s:
        return None
    try:
        x = float(s)
    except Exception:
        return None
    if not math.isfinite(x):
        return None
    return x


def haversine_m(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 6371008.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlambda / 2) ** 2
    return 2 * r * math.asin(min(1.0, math.sqrt(a)))


def distance_band(d: Optional[float]) -> str:
    if d is None:
        return "NO_COORD"
    if d <= 10: return "<=10m"
    if d <= 25: return "<=25m"
    if d <= 50: return "<=50m"
    if d <= 100: return "<=100m"
    if d <= 250: return "<=250m"
    if d <= 500: return "<=500m"
    return ">500m"


def pick_existing(patterns: Sequence[str], base: Path) -> Optional[Path]:
    hits: List[Path] = []
    for pat in patterns:
        hits.extend(base.glob(pat))
    hits = [p for p in hits if p.is_file()]
    if not hits:
        return None
    # Prefer exact/default-looking names, then newest mtime.
    hits.sort(key=lambda p: ("(" in p.name, -p.stat().st_mtime))
    return hits[0]


def sniff_csv(path: Path):
    fh = path.open("r", encoding="utf-8-sig", newline="")
    sample = fh.read(131072)
    fh.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except Exception:
        # ANNCSU mass CSVs are normally semicolon-delimited.
        delim = ";"
    reader = csv.DictReader(fh, delimiter=delim)
    return fh, reader, delim


def load_repair_log(path: Path) -> List[dict]:
    by_id: Dict[str, dict] = {}
    malformed = 0
    with path.open("r", encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
            except Exception:
                malformed += 1
                continue
            if obj.get("mode") != "repair":
                continue
            pid = str(obj.get("pvp_id") or "").strip()
            if pid:
                by_id[pid] = obj
    rows = list(by_id.values())
    rows.sort(key=lambda x: int(x.get("pvp_id") or 0))
    print(f"[LOG] repair unici: {len(rows)} (righe JSON malformate ignorate: {malformed})")
    if not rows:
        raise SystemExit("Nessun record mode=repair trovato nel log.")
    return rows


def download_if_missing(path: Path, url: str) -> None:
    if path.exists() and path.stat().st_size > 0 and zipfile.is_zipfile(path):
        return
    print(f"[ISTAT] scarico elenco Comuni ufficiale: {url}")
    req = urllib.request.Request(url, headers={"User-Agent": "AsteEasy-ANNCSU/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r, path.open("wb") as out:
            while True:
                chunk = r.read(1024 * 1024)
                if not chunk:
                    break
                out.write(chunk)
    except Exception as e:
        if path.exists():
            try: path.unlink()
            except Exception: pass
        raise SystemExit(
            "Impossibile scaricare automaticamente l'elenco ISTAT.\n"
            f"Scaricalo manualmente da:\n{url}\n"
            f"e salvalo come:\n{path}\n\nErrore: {e}"
        )


def _xlsx_shared_strings(z: zipfile.ZipFile) -> List[str]:
    try:
        data = z.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    root = ET.fromstring(data)
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    out = []
    for si in root.findall("m:si", ns):
        text = "".join(t.text or "" for t in si.findall(".//m:t", ns))
        out.append(text)
    return out


def _xlsx_first_sheet_path(z: zipfile.ZipFile) -> str:
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    sheet = wb.find("m:sheets/m:sheet", ns)
    if sheet is None:
        raise RuntimeError("Nessun foglio XLSX")
    rid = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rns = {"p": "http://schemas.openxmlformats.org/package/2006/relationships"}
    target = None
    for rel in rels.findall("p:Relationship", rns):
        if rel.attrib.get("Id") == rid:
            target = rel.attrib.get("Target")
            break
    if not target:
        raise RuntimeError("Relazione foglio XLSX non trovata")
    target = target.lstrip("/")
    if not target.startswith("xl/"):
        target = "xl/" + target
    return target


def _col_index(cell_ref: str) -> int:
    letters = re.match(r"([A-Z]+)", cell_ref or "")
    if not letters:
        return 0
    n = 0
    for ch in letters.group(1):
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def read_xlsx_rows_stdlib(path: Path) -> Iterable[List[str]]:
    with zipfile.ZipFile(path) as z:
        shared = _xlsx_shared_strings(z)
        sheet_path = _xlsx_first_sheet_path(z)
        data = z.read(sheet_path)
    root = ET.fromstring(data)
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    for row in root.findall(".//m:sheetData/m:row", ns):
        vals: Dict[int, str] = {}
        max_i = -1
        for c in row.findall("m:c", ns):
            i = _col_index(c.attrib.get("r", "A1"))
            max_i = max(max_i, i)
            t = c.attrib.get("t")
            if t == "inlineStr":
                val = "".join(x.text or "" for x in c.findall(".//m:t", ns))
            else:
                v = c.find("m:v", ns)
                raw = "" if v is None or v.text is None else v.text
                if t == "s" and raw:
                    try: val = shared[int(raw)]
                    except Exception: val = raw
                else:
                    val = raw
            vals[i] = val
        if max_i >= 0:
            yield [vals.get(i, "") for i in range(max_i + 1)]


def read_istat_rows(path: Path) -> Iterable[List[str]]:
    try:
        import openpyxl  # type: ignore
    except Exception:
        yield from read_xlsx_rows_stdlib(path)
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(values_only=True):
        yield ["" if v is None else str(v) for v in row]


def load_istat_comuni(path: Path) -> List[dict]:
    rows = read_istat_rows(path)
    header: Optional[List[str]] = None
    idx = {}
    out: List[dict] = []
    for row in rows:
        norms = [norm_text(x) for x in row]
        if header is None:
            # Header detection robust to minor wording changes.
            if any("denominazione in italiano" in x for x in norms) and any("sigla automobilistica" in x for x in norms):
                header = row
                for i, h in enumerate(norms):
                    if "denominazione in italiano" in h:
                        idx["name"] = i
                    elif "sigla automobilistica" in h:
                        idx["prov"] = i
                    elif "codice comune formato alfanumerico" in h:
                        idx["code"] = i
                continue
            continue
        if not all(k in idx for k in ("name", "prov", "code")):
            raise SystemExit("Colonne ISTAT richieste non riconosciute nell'xlsx.")
        def cell(k: str) -> str:
            i = idx[k]
            return row[i].strip() if i < len(row) else ""
        name, prov, code = cell("name"), cell("prov"), cell("code")
        if not name or not code:
            continue
        code_digits = re.sub(r"\D", "", code)
        if not code_digits:
            continue
        code6 = code_digits.zfill(6)[-6:]
        out.append({"code": code6, "name": name, "prov": prov.upper().strip(), "name_norm": norm_name(name)})
    if not out:
        raise SystemExit("Elenco ISTAT vuoto/non riconosciuto.")
    print(f"[ISTAT] Comuni caricati: {len(out)}")
    return out


def resolve_municipalities(records: List[dict], comuni: List[dict]) -> Tuple[Dict[str, str], List[dict]]:
    by_name: Dict[str, List[dict]] = defaultdict(list)
    for c in comuni:
        by_name[c["name_norm"]].append(c)
    resolved: Dict[str, str] = {}
    unresolved: List[dict] = []
    for r in records:
        pid = str(r.get("pvp_id"))
        name_n = norm_name(r.get("comune"))
        prov = str(r.get("sigla_provincia") or "").upper().strip()
        cand = by_name.get(name_n, [])
        chosen = None
        if len(cand) == 1:
            chosen = cand[0]
        elif cand:
            exact_prov = [c for c in cand if c["prov"] == prov]
            if len(exact_prov) == 1:
                chosen = exact_prov[0]
        if chosen:
            resolved[pid] = chosen["code"]
        else:
            unresolved.append(r)
    print(f"[ISTAT] record con Comune risolto: {len(resolved)}/{len(records)}; irrisolti: {len(unresolved)}")
    return resolved, unresolved


def source_street_tokens(r: dict) -> Tuple[str, set]:
    s = norm_text(r.get("indirizzo"))
    comune = norm_text(r.get("comune"))
    cap = norm_text(r.get("cap"))
    prov = norm_text(r.get("sigla_provincia"))
    # Remove known tail/context components but do not globally remove numbers: they can be part of street names.
    if comune:
        s = re.sub(r"\b" + re.escape(comune) + r"\b", " ", s)
    for x in (cap, prov):
        if x:
            s = re.sub(r"\b" + re.escape(x) + r"\b", " ", s)
    # Remove expected civics as standalone tokens only.
    exp = [normalize_civic(x) for x in str(r.get("expected_house_numbers") or "").split("|") if x]
    for x in exp:
        if x.isdigit():
            s = re.sub(r"\b" + re.escape(x) + r"\b", " ", s)

    # v10.0.2: when the source contains an explicit road type (Via, Piazza,
    # Strada, ...), remove separate locality/frazione context from the street
    # matching text. This prevents e.g.
    #   "Via delle Gardenie, 6, Monte Caminetto"
    # from matching ANNCSU "VIA MONTE CAMINETTO" merely because the locality
    # is present in the full PVP address. If no explicit road type is present
    # (e.g. rural "Cascina Casa del Rosso 8"), keep the toponym available.
    raw_tokens = set(norm_text(r.get("indirizzo")).split())
    explicit_road = bool(raw_tokens & STREET_KIND_TOKENS)
    if explicit_road:
        try:
            loc_hints = extract_source_locality_hints(r)
        except Exception:
            loc_hints = []
        for h in loc_hints:
            hn = norm_text(h)
            if hn:
                s = re.sub(r"(?<![a-z0-9])" + re.escape(hn) + r"(?![a-z0-9])", " ", s)

    s = re.sub(r"\s+", " ", s).strip()
    return s, set(tokens(s))


def street_variants(row: dict) -> List[str]:
    vals = []
    for key in ("ODONIMO", "DIZIONE_LINGUA1", "DIZIONE_LINGUA2"):
        v = (row.get(key) or "").strip()
        if v:
            vals.append(v)
    return vals or [(row.get("ODONIMO") or "").strip()]


def street_score(src_norm: str, src_tokens: set, odonimo: str) -> float:
    """v2: mantiene il matcher v1 e aggiunge un rescue conservativo per alias."""
    st_norm = norm_text(odonimo)
    st_tokens = set(tokens(st_norm))
    if not st_norm:
        return 0.0
    if st_norm in src_norm:
        base = 1.0
    elif not st_tokens:
        base = difflib.SequenceMatcher(None, st_norm, src_norm).ratio() * 0.65
    else:
        inter = st_tokens & src_tokens
        coverage = len(inter) / len(st_tokens)
        if coverage == 1.0:
            base = 0.96
        else:
            precision = len(inter) / max(1, len(src_tokens))
            f1 = 0.0 if not inter else (2 * coverage * precision / (coverage + precision))
            seq = difflib.SequenceMatcher(None, st_norm, src_norm).ratio()
            base = max(0.72 * coverage + 0.18 * f1 + 0.10 * seq, seq * 0.72)
    return max(base, alias_street_score(src_norm, odonimo))


def prepare_record_indexes(records: List[dict], resolved: Dict[str, str]):
    by_code: Dict[str, List[int]] = defaultdict(list)
    token_index: Dict[str, Dict[str, set]] = defaultdict(lambda: defaultdict(set))
    source_info = {}
    for i, r in enumerate(records):
        pid = str(r.get("pvp_id"))
        code = resolved.get(pid)
        if not code:
            continue
        src_norm, src_toks = source_street_tokens(r)
        source_info[pid] = (src_norm, src_toks)
        by_code[code].append(i)
        for t in src_toks:
            token_index[code][t].add(i)
    return by_code, token_index, source_info


def scan_strad(path: Path, records: List[dict], resolved: Dict[str, str], topn: int = 3):
    by_code, token_index, source_info = prepare_record_indexes(records, resolved)
    target_codes = set(by_code)
    top: Dict[str, List[Tuple[float, str, str, str]]] = defaultdict(list)
    fh, reader, delim = sniff_csv(path)
    print(f"[STRAD] delimiter={delim!r}; Comuni target={len(target_codes)}")
    n = kept = 0
    try:
        for row in reader:
            n += 1
            if n % 500000 == 0:
                print(f"[STRAD] lette {n:,} righe; utili {kept:,}")
            code = str(row.get("CODICE_ISTAT") or "").strip().zfill(6)
            if code not in target_codes:
                continue
            prog = str(row.get("PROGRESSIVO_NAZIONALE") or "").strip()
            odo = str(row.get("ODONIMO") or "").strip()
            if not prog or not odo:
                continue
            variants = street_variants(row)
            street_toks = set()
            for v in variants:
                street_toks |= set(tokens(v))
            candidate_idx = set()
            for t in street_toks:
                candidate_idx |= token_index[code].get(t, set())
            if not candidate_idx:
                continue
            kept += 1
            for i in candidate_idx:
                r = records[i]
                pid = str(r.get("pvp_id"))
                src_norm, src_toks = source_info[pid]
                best_v = max(variants, key=lambda v: street_score(src_norm, src_toks, v))
                sc = street_score(src_norm, src_toks, best_v)
                if sc < 0.35:
                    continue
                loc = str(row.get("LOCALITA'") or "").strip()
                arr = top[pid]
                arr.append((sc, prog, odo, loc))
                arr.sort(key=lambda x: x[0], reverse=True)
                # unique progressivo, top N
                dedup = []
                seen = set()
                for item in arr:
                    if item[1] in seen:
                        continue
                    seen.add(item[1]); dedup.append(item)
                    if len(dedup) >= topn:
                        break
                top[pid] = dedup
    finally:
        fh.close()
    matched = sum(1 for pid in resolved if top.get(pid))
    print(f"[STRAD] fine: {n:,} righe; record con almeno una strada candidata: {matched}/{len(resolved)}")
    return top


def expected_civics(r: dict) -> set:
    vals = set()
    raw = str(r.get("expected_house_numbers") or "")
    for x in raw.split("|"):
        c = normalize_civic(x)
        if c:
            vals.add(c)
    # Add obvious number+letter forms directly visible in source, e.g. 12/A, 12A.
    src = norm_text(r.get("indirizzo"))
    for m in re.finditer(r"\b(\d{1,5})\s*[/\-]?\s*([a-z]{1,3})\b", src):
        vals.add(normalize_civic(m.group(1) + m.group(2)))
    return vals


def civic_relation(expected: set, actual: str) -> str:
    if not expected:
        return "not_expected"
    a = normalize_civic(actual)
    if not a:
        return "missing"
    if a in expected:
        return "exact"
    abase = re.match(r"(\d+)", a)
    if abase:
        b = abase.group(1)
        for e in expected:
            eb = re.match(r"(\d+)", e)
            if eb and eb.group(1) == b:
                return "variant"
    return "mismatch"


def scan_indir(path: Path, wanted_progs: set):
    accesses: Dict[str, List[Tuple[str, Optional[float], Optional[float], str]]] = defaultdict(list)
    fh, reader, delim = sniff_csv(path)
    print(f"[INDIR] delimiter={delim!r}; strade candidate da trattenere={len(wanted_progs):,}")
    n = kept = coords = 0
    try:
        for row in reader:
            n += 1
            if n % 1000000 == 0:
                print(f"[INDIR] lette {n:,} righe; accessi utili {kept:,}; con coordinate {coords:,}")
            prog = str(row.get("PROGRESSIVO_NAZIONALE") or "").strip()
            if prog not in wanted_progs:
                continue
            civ = normalize_civic(str(row.get("CIVICO") or "") + str(row.get("ESPONENTE") or ""))
            lon = safe_float(row.get("COORD_X_COMUNE"))
            lat = safe_float(row.get("COORD_Y_COMUNE"))
            metodo = str(row.get("METODO") or "").strip()
            kept += 1
            if lat is not None and lon is not None:
                coords += 1
            accesses[prog].append((civ, lat, lon, metodo))
    finally:
        fh.close()
    print(f"[INDIR] fine: {n:,} righe; accessi utili {kept:,}; con coordinate {coords:,}")
    return accesses


def google_bucket(r: dict) -> str:
    action = str(r.get("action") or "")
    status = str(r.get("status") or "")
    if action == "corrected_from_pvp":
        return "CORRECTED_FROM_PVP"
    if status == "PVP_OK_CIVICO":
        return "PVP_OK_CIVICO"
    if status == "PVP_OK_STRADA":
        return "PVP_OK_STRADA"
    if status == "PVP_SUSPECT":
        return "PVP_SUSPECT"
    if status == "PVP_UNCERTAIN":
        return "PVP_UNCERTAIN"
    return status or "OTHER"


def evaluate_record(r: dict, code: Optional[str], street_cands, accesses) -> dict:
    out = {
        "pvp_id": str(r.get("pvp_id") or ""),
        "google_bucket": google_bucket(r),
        "google_status": str(r.get("status") or ""),
        "google_reason": str(r.get("reason") or ""),
        "google_action": str(r.get("action") or ""),
        "indirizzo": str(r.get("indirizzo") or ""),
        "comune": str(r.get("comune") or ""),
        "sigla_provincia": str(r.get("sigla_provincia") or ""),
        "cap": str(r.get("cap") or ""),
        "codice_istat": code or "",
        "expected_house_numbers": str(r.get("expected_house_numbers") or ""),
        "anncsu_odonimo": "",
        "anncsu_localita": "",
        "street_score": "",
        "civic_relation": "",
        "anncsu_civic": "",
        "anncsu_metodo": "",
        "anncsu_lat": "",
        "anncsu_lon": "",
        "distance_pvp_m": "",
        "distance_google_corrected_m": "",
        "distance_band": "",
        "local_signal": "",
    }
    if not code:
        out["local_signal"] = "COMUNE_NON_RISOLTO"
        return out
    if not street_cands:
        out["local_signal"] = "STRADA_NON_TROVATA"
        return out

    exp = expected_civics(r)
    old_lat, old_lon = safe_float(r.get("old_lat")), safe_float(r.get("old_lon"))
    sug_lat, sug_lon = safe_float(r.get("suggested_lat")), safe_float(r.get("suggested_lon"))

    choices = []
    for sc, prog, odo, loc in street_cands:
        acc = accesses.get(prog, [])
        if exp:
            exact = [a for a in acc if civic_relation(exp, a[0]) == "exact"]
            variant = [a for a in acc if civic_relation(exp, a[0]) == "variant"]
            pool = exact or variant
            rel = "exact" if exact else ("variant" if variant else "not_found")
        else:
            pool = acc
            rel = "not_expected"
        if not pool:
            # Keep street-only candidate too.
            choices.append((0 if exp else 1, sc, None, prog, odo, loc, rel))
            continue
        # Prefer georeferenced access nearest to original PVP coordinate; otherwise first access.
        best = None
        best_d = None
        for a in pool:
            civ, lat, lon, metodo = a
            d = None
            if old_lat is not None and old_lon is not None and lat is not None and lon is not None:
                d = haversine_m(old_lat, old_lon, lat, lon)
            key = (d is None, d if d is not None else 1e99)
            if best is None or key < best[0]:
                best = (key, a)
                best_d = d
        civic_rank = 3 if rel == "exact" else (2 if rel == "variant" else 1)
        choices.append((civic_rank, sc, best[1] if best else None, prog, odo, loc, rel))

    choices.sort(key=lambda x: (x[0], x[1]), reverse=True)
    civic_rank, sc, access, prog, odo, loc, rel = choices[0]
    out["anncsu_odonimo"] = odo
    out["anncsu_localita"] = loc
    out["street_score"] = round(sc, 4)
    out["civic_relation"] = rel

    d_old = None
    if access:
        civ, lat, lon, metodo = access
        out["anncsu_civic"] = civ
        out["anncsu_metodo"] = metodo
        out["anncsu_lat"] = "" if lat is None else lat
        out["anncsu_lon"] = "" if lon is None else lon
        if old_lat is not None and old_lon is not None and lat is not None and lon is not None:
            d_old = haversine_m(old_lat, old_lon, lat, lon)
            out["distance_pvp_m"] = round(d_old, 2)
            out["distance_band"] = distance_band(d_old)
        if sug_lat is not None and sug_lon is not None and lat is not None and lon is not None:
            out["distance_google_corrected_m"] = round(haversine_m(sug_lat, sug_lon, lat, lon), 2)

    # Conservative signal labels only for benchmarking; they DO NOT alter DB.
    if sc < 0.72:
        signal = "STREET_MATCH_DEBOLE"
    elif exp and rel == "exact":
        if d_old is None:
            signal = "CIVICO_UFFICIALE_SENZA_COORD"
        elif d_old <= 25:
            signal = "LOCAL_OK_CIVICO_25M"
        elif d_old <= 50:
            signal = "LOCAL_OK_CIVICO_50M"
        elif d_old <= 100:
            signal = "LOCAL_OK_CIVICO_100M"
        elif d_old > 250:
            signal = "LOCAL_CIVICO_LONTANO_250M"
        else:
            signal = "LOCAL_CIVICO_INTERMEDIO"
    elif exp and rel == "variant":
        signal = "CIVICO_VARIANTE"
    elif exp:
        signal = "CIVICO_NON_TROVATO"
    else:
        if d_old is None:
            signal = "STRADA_UFFICIALE_SENZA_COORD"
        elif d_old <= 50:
            signal = "LOCAL_OK_STRADA_50M"
        elif d_old <= 100:
            signal = "LOCAL_OK_STRADA_100M"
        elif d_old > 250:
            signal = "LOCAL_STRADA_LONTANA_250M"
        else:
            signal = "LOCAL_STRADA_INTERMEDIA"
    out["local_signal"] = signal
    return out


def build_summary(rows: List[dict], unresolved_count: int) -> dict:
    total = len(rows)
    by_google = Counter(r["google_bucket"] for r in rows)
    by_signal = Counter(r["local_signal"] for r in rows)
    by_band = Counter(r["distance_band"] or "NO_DISTANCE" for r in rows)

    def filt(bucket):
        return [r for r in rows if r["google_bucket"] == bucket]

    def count_if(arr, pred):
        return sum(1 for r in arr if pred(r))

    thresholds = [10, 25, 50, 100, 250]
    bench = {}
    for bucket in ["PVP_OK_CIVICO", "PVP_OK_STRADA", "CORRECTED_FROM_PVP", "PVP_SUSPECT", "PVP_UNCERTAIN"]:
        arr = filt(bucket)
        if not arr:
            continue
        item = {"n": len(arr)}
        item["street_score_ge_0_72_v2"] = count_if(arr, lambda r: isinstance(r["street_score"], (int, float)) and r["street_score"] >= 0.72)
        item["civic_exact"] = count_if(arr, lambda r: r["civic_relation"] == "exact")
        for t in thresholds:
            item[f"distance_pvp_le_{t}m"] = count_if(arr, lambda r, t=t: isinstance(r["distance_pvp_m"], (int, float)) and r["distance_pvp_m"] <= t)
            item[f"distance_pvp_gt_{t}m"] = count_if(arr, lambda r, t=t: isinstance(r["distance_pvp_m"], (int, float)) and r["distance_pvp_m"] > t)
            item[f"distance_google_corrected_le_{t}m"] = count_if(arr, lambda r, t=t: isinstance(r["distance_google_corrected_m"], (int, float)) and r["distance_google_corrected_m"] <= t)
        bench[bucket] = item

    return {
        "version": VERSION,
        "generated_at": datetime.now().astimezone().isoformat(),
        "records": total,
        "municipality_unresolved": unresolved_count,
        "by_google_bucket": dict(by_google),
        "by_local_signal": dict(by_signal),
        "by_distance_band": dict(by_band),
        "benchmark": bench,
        "notes": [
            "Nessun dato Supabase viene modificato.",
            "Nessuna API Google viene chiamata.",
            "Le soglie locali sono esplorative: vanno calibrate confrontandole con i bucket Google.",
            "v2 aggiunge normalizzazione alias: iniziali, numeri romani/date e nomi ufficiali estesi.",
            "Per CORRECTED_FROM_PVP, distance_pvp_m misura la distanza ANNCSU dalle vecchie coordinate PVP; distance_google_corrected_m dalle coordinate corrette da Google.",
        ],
    }


def main() -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="Benchmark ANNCSU locale sui record già verificati con Google")
    ap.add_argument("--tools-dir", type=Path, default=here)
    ap.add_argument("--strad", type=Path)
    ap.add_argument("--indir", type=Path)
    ap.add_argument("--log", type=Path)
    ap.add_argument("--istat-xlsx", type=Path)
    ap.add_argument("--top-streets", type=int, default=3)
    args = ap.parse_args()

    base = args.tools_dir.resolve()
    strad = args.strad or pick_existing(["STRAD_ITA_*.csv"], base)
    indir = args.indir or pick_existing(["INDIR_ITA_*.csv"], base)
    log = args.log or pick_existing(["google_geocode_log_v8.jsonl", "google_geocode_log_v8*.jsonl"], base)
    istat = args.istat_xlsx or (base / "Elenco-comuni-italiani.xlsx")

    for label, path in (("STRAD", strad), ("INDIR", indir), ("LOG", log)):
        if path is None or not Path(path).exists():
            raise SystemExit(f"File {label} non trovato in {base}")
    strad, indir, log, istat = Path(strad), Path(indir), Path(log), Path(istat)

    print("=" * 72)
    print(f"Aste Easy - ANNCSU benchmark v{VERSION}")
    print(f"tools : {base}")
    print(f"STRAD : {strad.name}")
    print(f"INDIR : {indir.name}")
    print(f"LOG   : {log.name}")
    print("Modalità READ-ONLY: nessun update Supabase, nessuna chiamata Google.")
    print("=" * 72)

    t0 = time.time()
    records = load_repair_log(log)
    download_if_missing(istat, ISTAT_XLSX_URL)
    comuni = load_istat_comuni(istat)
    resolved, unresolved = resolve_municipalities(records, comuni)

    top = scan_strad(strad, records, resolved, max(1, args.top_streets))
    wanted = {prog for arr in top.values() for _, prog, _, _ in arr}
    print(f"[MATCH] progressivi ANNCSU candidati complessivi: {len(wanted):,}")
    accesses = scan_indir(indir, wanted)

    result_rows = []
    for r in records:
        pid = str(r.get("pvp_id"))
        result_rows.append(evaluate_record(r, resolved.get(pid), top.get(pid, []), accesses))

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_csv = base / f"anncsu_benchmark_v2_{stamp}.csv"
    out_json = base / f"anncsu_benchmark_summary_v2_{stamp}.json"
    out_unresolved = base / f"anncsu_comuni_irrisolti_v2_{stamp}.csv"

    fields = list(result_rows[0].keys())
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader(); w.writerows(result_rows)

    summary = build_summary(result_rows, len(unresolved))
    summary["files"] = {"strad": strad.name, "indir": indir.name, "google_log": log.name, "istat": istat.name}
    summary["elapsed_seconds"] = round(time.time() - t0, 2)
    out_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    with out_unresolved.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["pvp_id", "comune", "sigla_provincia", "cap", "indirizzo"])
        w.writeheader()
        for r in unresolved:
            w.writerow({k: r.get(k, "") for k in ["pvp_id", "comune", "sigla_provincia", "cap", "indirizzo"]})

    # Diagnostic-only report: Google corrections that materially disagree with ANNCSU.
    out_conflicts = base / f"anncsu_google_conflicts_v2_{stamp}.csv"
    conflict_rows = []
    for r in result_rows:
        if r.get("google_bucket") != "CORRECTED_FROM_PVP":
            continue
        d_old = r.get("distance_pvp_m")
        d_new = r.get("distance_google_corrected_m")
        if not isinstance(d_old, (int, float)) or not isinstance(d_new, (int, float)):
            continue
        signal = ""
        if d_old <= 25 and d_new > 50:
            signal = "PVP_VICINO_ANNCSU_GOOGLE_LONTANO"
        elif d_new <= 25 and d_old > 50:
            signal = "GOOGLE_VICINO_ANNCSU_PVP_LONTANO"
        elif d_old <= 25 and d_new <= 25:
            signal = "ENTRAMBI_ENTRO_25M"
        elif d_old > 100 and d_new > 100:
            signal = "ENTRAMBI_OLTRE_100M"
        elif d_new + 25 < d_old:
            signal = "GOOGLE_PIU_VICINO_ANNCSU"
        elif d_old + 25 < d_new:
            signal = "PVP_PIU_VICINO_ANNCSU"
        if signal:
            rr = dict(r)
            rr["comparison_signal"] = signal
            conflict_rows.append(rr)
    if conflict_rows:
        with out_conflicts.open("w", encoding="utf-8-sig", newline="") as f:
            fields_conf = list(conflict_rows[0].keys())
            w = csv.DictWriter(f, fieldnames=fields_conf)
            w.writeheader(); w.writerows(conflict_rows)
    else:
        out_conflicts.write_text("", encoding="utf-8")

    # Strong zero-Google diagnostic potential (still benchmark-only).
    strong_zero = sum(
        1 for r in result_rows
        if isinstance(r.get("street_score"), (int, float)) and r["street_score"] >= 0.72
        and (
            (r.get("civic_relation") == "exact" and isinstance(r.get("distance_pvp_m"), (int, float)) and r["distance_pvp_m"] <= 25)
            or
            (r.get("civic_relation") == "not_expected" and isinstance(r.get("distance_pvp_m"), (int, float)) and r["distance_pvp_m"] <= 50)
        )
    )
    summary["strong_zero_google_candidates"] = strong_zero
    summary["strong_zero_google_pct"] = round(100.0 * strong_zero / len(result_rows), 2) if result_rows else 0.0
    summary["google_anncsu_conflict_rows"] = len(conflict_rows)
    summary["files"]["conflicts"] = out_conflicts.name
    out_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n=== RISULTATO ===")
    print(f"CSV dettagli : {out_csv}")
    print(f"JSON summary : {out_json}")
    print(f"Comuni dubbi : {out_unresolved}")
    print(f"Conflitti G/A: {out_conflicts}")
    print("\nGoogle bucket:")
    for k, v in summary["by_google_bucket"].items():
        print(f"  {k:24s} {v:6d}")
    print("\nSegnali ANNCSU:")
    for k, v in sorted(summary["by_local_signal"].items(), key=lambda kv: (-kv[1], kv[0])):
        print(f"  {k:30s} {v:6d}")
    print("\nNon modificare ancora il repair sulla base di queste soglie: prima leggiamo il summary e calibriamo sui risultati Google.")
    return 0




# ============================================================================
# v3 overrides: locality-aware, ambiguity-aware, production-decision simulation
# ============================================================================
VERSION = "3.0"

# Historical/merged municipality aliases encountered in the benchmark.  The
# resolver remains conservative: an alias is accepted only when it uniquely
# identifies a current ISTAT municipality. Province inconsistencies are exposed
# separately and prevent LOCAL_ACCEPT.
MUNICIPALITY_ALIASES = {
    "veruno": "gattico veruno",
    "tavarnelle val di pesa": "barberino tavarnelle",
    "montebello ionico": "montebello jonico",
    "presicce": "presicce acquarica",
    "tonengo": "moransengo tonengo",
    "vighizzolo d este": "borgo veneto",
}

LOCALITY_PREFIX_RE = re.compile(
    r"(?:^|[,;])\s*(?:localita|loc|frazione|fraz|borgata|borgo)\.?\s+([^,;]+)", re.I
)
STREET_KIND_TOKENS = {
    "via", "viale", "piazza", "corso", "strada", "vicolo", "vico", "largo",
    "salita", "traversa", "contrada", "cda", "ss", "sp", "sr", "sc"
}


def resolve_municipalities_v3(records: List[dict], comuni: List[dict]):
    by_name: Dict[str, List[dict]] = defaultdict(list)
    for c in comuni:
        by_name[c["name_norm"]].append(c)
    resolved: Dict[str, str] = {}
    meta: Dict[str, dict] = {}
    unresolved: List[dict] = []

    for r in records:
        pid = str(r.get("pvp_id"))
        source_name = norm_name(r.get("comune"))
        source_prov = str(r.get("sigla_provincia") or "").upper().strip()
        target_name = source_name
        resolution = "EXACT"
        cand = by_name.get(target_name, [])
        if not cand and source_name in MUNICIPALITY_ALIASES:
            target_name = MUNICIPALITY_ALIASES[source_name]
            cand = by_name.get(target_name, [])
            resolution = "ALIAS"

        chosen = None
        if len(cand) == 1:
            chosen = cand[0]
        elif cand:
            exact_prov = [c for c in cand if c["prov"] == source_prov]
            if len(exact_prov) == 1:
                chosen = exact_prov[0]

        if chosen:
            resolved[pid] = chosen["code"]
            prov_rel = "MATCH" if (not source_prov or source_prov == chosen["prov"]) else "MISMATCH"
            meta[pid] = {
                "municipality_resolution": resolution,
                "istat_comune": chosen["name"],
                "istat_province": chosen["prov"],
                "province_relation": prov_rel,
            }
        else:
            unresolved.append(r)
            meta[pid] = {
                "municipality_resolution": "UNRESOLVED",
                "istat_comune": "",
                "istat_province": "",
                "province_relation": "UNKNOWN",
            }
    print(f"[ISTAT v3] record con Comune risolto: {len(resolved)}/{len(records)}; irrisolti: {len(unresolved)}")
    return resolved, meta, unresolved


def extract_source_locality_hints(r: dict) -> List[str]:
    raw = str(r.get("indirizzo") or "")
    hints: List[str] = []

    # Explicit locality/frazione markers have highest semantic value.
    for m in LOCALITY_PREFIX_RE.finditer(raw.replace("’", "'")):
        h = norm_text(m.group(1))
        h = re.sub(r"\b\d{5}\b.*$", "", h).strip()
        if h and len(h.split()) <= 6:
            hints.append(h)

    # Comma-separated context components such as "Sant'Angelo" or "Mariotto".
    comune = norm_text(r.get("comune"))
    prov = norm_text(r.get("sigla_provincia"))
    cap = norm_text(r.get("cap"))
    for part in re.split(r"[,;]", raw):
        p = norm_text(part)
        if not p or p == comune or p == prov or p == cap:
            continue
        if comune and comune in p:
            continue
        if cap and cap in p:
            continue
        ts = p.split()
        if any(t in STREET_KIND_TOKENS for t in ts):
            continue
        if any(re.fullmatch(r"\d+[a-z]?", t) for t in ts):
            continue
        if 1 <= len(ts) <= 5 and any(len(t) >= 3 for t in ts):
            hints.append(p)

    # Deduplicate while preserving order.
    out, seen = [], set()
    for h in hints:
        h = norm_text(h)
        if h and h not in seen:
            seen.add(h); out.append(h)
    return out[:6]


def locality_similarity(hints: List[str], ann_localita: str) -> Optional[float]:
    ann = norm_text(ann_localita)
    if not ann:
        return None
    if not hints:
        return None
    at = set(tokens(ann, drop_stop=False))
    best = 0.0
    for h in hints:
        hn = norm_text(h)
        if not hn:
            continue
        if hn == ann:
            return 1.0
        if hn in ann or ann in hn:
            best = max(best, 0.95)
        ht = set(tokens(hn, drop_stop=False))
        inter = len(ht & at)
        if inter:
            cov = inter / max(1, min(len(ht), len(at)))
            seq = difflib.SequenceMatcher(None, hn, ann).ratio()
            best = max(best, 0.72 * cov + 0.28 * seq)
        else:
            best = max(best, difflib.SequenceMatcher(None, hn, ann).ratio() * 0.45)
    return min(1.0, best)


def _coord_spread_m(access_list) -> Optional[float]:
    pts = [(a[1], a[2]) for a in access_list if a[1] is not None and a[2] is not None]
    if len(pts) < 2:
        return 0.0 if pts else None
    # Bounding diagonal is a safe, cheap upper-bound-like indicator for ambiguity.
    lats = [p[0] for p in pts]; lons = [p[1] for p in pts]
    return haversine_m(min(lats), min(lons), max(lats), max(lons))


def _min_distance_to_accesses(lat: Optional[float], lon: Optional[float], acc_list):
    if lat is None or lon is None:
        return None, None
    best = None
    for a in acc_list:
        if a[1] is None or a[2] is None:
            continue
        d = haversine_m(lat, lon, a[1], a[2])
        if best is None or d < best[0]:
            best = (d, a)
    return best if best is not None else (None, None)


def evaluate_record_v3(r: dict, code: Optional[str], municipality_meta: dict, street_cands, accesses) -> dict:
    hints = extract_source_locality_hints(r)
    out = {
        "pvp_id": str(r.get("pvp_id") or ""),
        "google_bucket": google_bucket(r),
        "google_status": str(r.get("status") or ""),
        "google_reason": str(r.get("reason") or ""),
        "google_action": str(r.get("action") or ""),
        "indirizzo": str(r.get("indirizzo") or ""),
        "comune": str(r.get("comune") or ""),
        "sigla_provincia": str(r.get("sigla_provincia") or ""),
        "cap": str(r.get("cap") or ""),
        "codice_istat": code or "",
        "municipality_resolution": municipality_meta.get("municipality_resolution", "UNRESOLVED"),
        "istat_comune": municipality_meta.get("istat_comune", ""),
        "istat_province": municipality_meta.get("istat_province", ""),
        "province_relation": municipality_meta.get("province_relation", "UNKNOWN"),
        "source_locality_hints": "|".join(hints),
        "expected_house_numbers": str(r.get("expected_house_numbers") or ""),
        "anncsu_odonimo": "",
        "anncsu_localita": "",
        "locality_score": "",
        "street_score": "",
        "civic_relation": "",
        "anncsu_civic": "",
        "anncsu_metodo": "",
        "anncsu_lat": "",
        "anncsu_lon": "",
        "matching_street_candidates": 0,
        "matching_access_candidates": 0,
        "candidate_coord_spread_m": "",
        "ambiguity": "",
        "ambiguity_reason": "",
        "distance_pvp_m": "",
        "distance_google_corrected_m": "",
        "distance_band": "",
        "local_signal": "",
        "anncsu_decision": "",
        "decision_reason": "",
    }

    if not code:
        out["local_signal"] = "COMUNE_NON_RISOLTO"
        out["anncsu_decision"] = "GOOGLE_NEEDED"
        out["decision_reason"] = "municipality_unresolved"
        return out
    if municipality_meta.get("province_relation") == "MISMATCH":
        # Continue the ANNCSU analysis for diagnostics, but never auto-accept.
        province_bad = True
    else:
        province_bad = False
    if not street_cands:
        out["local_signal"] = "STRADA_NON_TROVATA"
        out["anncsu_decision"] = "GOOGLE_NEEDED"
        out["decision_reason"] = "street_not_found"
        return out

    exp = expected_civics(r)
    old_lat, old_lon = safe_float(r.get("old_lat")), safe_float(r.get("old_lon"))
    sug_lat, sug_lon = safe_float(r.get("suggested_lat")), safe_float(r.get("suggested_lon"))

    candidates = []
    for sc, prog, odo, loc in street_cands:
        acc = accesses.get(prog, [])
        loc_sc = locality_similarity(hints, loc)
        if exp:
            exact = [a for a in acc if civic_relation(exp, a[0]) == "exact"]
            variant = [a for a in acc if civic_relation(exp, a[0]) == "variant"]
            pool = exact or variant
            rel = "exact" if exact else ("variant" if variant else "not_found")
        else:
            pool = acc
            rel = "not_expected"
        rel_rank = {"exact": 4, "variant": 3, "not_expected": 2, "not_found": 1}.get(rel, 0)
        # Selection deliberately excludes PVP distance. Locality is only a positive
        # tie-breaker; missing ANNCSU locality is not treated as a mismatch.
        loc_rank = 2 if isinstance(loc_sc, (int, float)) and loc_sc >= 0.84 else (1 if loc_sc is None else 0)
        candidates.append({
            "sc": sc, "prog": prog, "odo": odo, "loc": loc, "loc_sc": loc_sc,
            "acc": acc, "pool": pool, "rel": rel, "rel_rank": rel_rank, "loc_rank": loc_rank,
        })

    # Best semantic relation first, locality second, street similarity third.
    candidates.sort(key=lambda c: (c["rel_rank"], c["loc_rank"], c["sc"]), reverse=True)
    best = candidates[0]

    # Consider semantically competing candidates. A locality match can safely
    # disambiguate; otherwise spatially separated duplicate streets/civics remain ambiguous.
    comparable = [c for c in candidates if c["rel_rank"] == best["rel_rank"] and c["sc"] >= max(0.72, best["sc"] - 0.04)]
    locality_winners = [c for c in comparable if isinstance(c["loc_sc"], (int, float)) and c["loc_sc"] >= 0.84]
    if locality_winners:
        comparable = locality_winners
        comparable.sort(key=lambda c: (c["loc_sc"], c["sc"]), reverse=True)
        best = comparable[0]

    out["matching_street_candidates"] = len(comparable)

    # Ambiguity across multiple ANNCSU street progressivi: inspect their georeferenced
    # pools without using PVP distance to choose among them.
    ambiguity = False
    ambiguity_reason = ""
    if len({c["prog"] for c in comparable}) > 1:
        reps = []
        for c in comparable:
            pts = [(a[1], a[2]) for a in c["pool"] if a[1] is not None and a[2] is not None]
            if pts:
                reps.append((c["prog"], pts[0][0], pts[0][1]))
        if len(reps) < len({c["prog"] for c in comparable}):
            ambiguity = True; ambiguity_reason = "multiple_streets_without_complete_coords"
        else:
            max_sep = 0.0
            for i in range(len(reps)):
                for j in range(i + 1, len(reps)):
                    max_sep = max(max_sep, haversine_m(reps[i][1], reps[i][2], reps[j][1], reps[j][2]))
            if max_sep > 150:
                ambiguity = True; ambiguity_reason = f"multiple_streets_separated_{round(max_sep)}m"

    pool = best["pool"]
    out["matching_access_candidates"] = len(pool)
    spread = _coord_spread_m(pool)
    if spread is not None:
        out["candidate_coord_spread_m"] = round(spread, 2)
    if len(pool) > 1 and spread is not None and spread > 150 and not ambiguity:
        ambiguity = True
        ambiguity_reason = f"same_address_accesses_separated_{round(spread)}m"

    out["ambiguity"] = "YES" if ambiguity else "NO"
    out["ambiguity_reason"] = ambiguity_reason
    out["anncsu_odonimo"] = best["odo"]
    out["anncsu_localita"] = best["loc"]
    out["locality_score"] = "" if best["loc_sc"] is None else round(best["loc_sc"], 4)
    out["street_score"] = round(best["sc"], 4)
    out["civic_relation"] = best["rel"]

    # Distances are to the set of official access points belonging to the selected,
    # semantically unambiguous address/street; nearest-point distance is legitimate here.
    d_old, best_old_access = _min_distance_to_accesses(old_lat, old_lon, pool)
    d_new, best_new_access = _min_distance_to_accesses(sug_lat, sug_lon, pool)
    report_access = best_old_access or best_new_access or (pool[0] if pool else None)
    if report_access:
        civ, lat, lon, metodo = report_access
        out["anncsu_civic"] = civ
        out["anncsu_metodo"] = metodo
        out["anncsu_lat"] = "" if lat is None else lat
        out["anncsu_lon"] = "" if lon is None else lon
    if d_old is not None:
        out["distance_pvp_m"] = round(d_old, 2)
        out["distance_band"] = distance_band(d_old)
    if d_new is not None:
        out["distance_google_corrected_m"] = round(d_new, 2)

    sc = best["sc"]
    rel = best["rel"]
    # Diagnostic local signal, compatible with v2 labels where possible.
    if sc < 0.72:
        signal = "STREET_MATCH_DEBOLE"
    elif ambiguity:
        signal = "ANNCSU_AMBIGUO"
    elif exp and rel == "exact":
        if d_old is None: signal = "CIVICO_UFFICIALE_SENZA_COORD"
        elif d_old <= 25: signal = "LOCAL_OK_CIVICO_25M"
        elif d_old <= 50: signal = "LOCAL_OK_CIVICO_50M"
        elif d_old <= 100: signal = "LOCAL_OK_CIVICO_100M"
        elif d_old > 250: signal = "LOCAL_CIVICO_LONTANO_250M"
        else: signal = "LOCAL_CIVICO_INTERMEDIO"
    elif exp and rel == "variant":
        signal = "CIVICO_VARIANTE"
    elif exp:
        signal = "CIVICO_NON_TROVATO"
    else:
        if d_old is None: signal = "STRADA_UFFICIALE_SENZA_COORD"
        elif d_old <= 50: signal = "LOCAL_OK_STRADA_50M"
        elif d_old <= 100: signal = "LOCAL_OK_STRADA_100M"
        elif d_old > 250: signal = "LOCAL_STRADA_LONTANA_250M"
        else: signal = "LOCAL_STRADA_INTERMEDIA"
    out["local_signal"] = signal

    # Production-like decision simulation. Conservative by design.
    decision = "GOOGLE_NEEDED"
    reason = "insufficient_local_evidence"
    if province_bad:
        decision, reason = "LOCAL_REVIEW", "source_province_disagrees_with_istat"
    elif ambiguity:
        decision, reason = "GOOGLE_NEEDED", "anncsu_ambiguous"
    elif sc < 0.72:
        decision, reason = "GOOGLE_NEEDED", "weak_street_match"
    elif exp and rel == "exact" and d_old is not None and d_old <= 25:
        decision, reason = "LOCAL_ACCEPT", "exact_civic_unambiguous_within_25m"
    elif not exp and rel == "not_expected" and d_old is not None and d_old <= 50:
        decision, reason = "LOCAL_ACCEPT", "unambiguous_street_within_50m"
    elif exp and rel == "exact" and d_old is None:
        decision, reason = "LOCAL_REVIEW", "official_exact_civic_without_coordinates"
    elif not exp and d_old is None:
        decision, reason = "LOCAL_REVIEW", "official_street_without_coordinates"
    elif exp and rel == "variant":
        decision, reason = "LOCAL_REVIEW", "official_civic_variant"
    elif exp and rel == "exact" and d_old is not None and d_old <= 100:
        decision, reason = "LOCAL_REVIEW", "exact_civic_between_25m_and_100m"
    elif not exp and d_old is not None and d_old <= 100:
        decision, reason = "LOCAL_REVIEW", "street_between_50m_and_100m"
    elif exp and rel == "exact" and d_old is not None and d_old > 100:
        decision, reason = "GOOGLE_NEEDED", "exact_official_civic_far_from_pvp"
    elif exp and rel == "not_found":
        decision, reason = "GOOGLE_NEEDED", "civic_not_found_in_anncsu"

    out["anncsu_decision"] = decision
    out["decision_reason"] = reason
    return out


def build_summary_v3(rows: List[dict], unresolved_count: int) -> dict:
    total = len(rows)
    by_google = Counter(r["google_bucket"] for r in rows)
    by_signal = Counter(r["local_signal"] for r in rows)
    by_decision = Counter(r["anncsu_decision"] for r in rows)
    by_reason = Counter(r["decision_reason"] for r in rows)
    by_ambiguity = Counter(r["ambiguity"] or "N/A" for r in rows)
    by_muni_res = Counter(r["municipality_resolution"] for r in rows)
    province_rel = Counter(r["province_relation"] for r in rows)

    matrix = {}
    for dec in ["LOCAL_ACCEPT", "LOCAL_REVIEW", "GOOGLE_NEEDED"]:
        matrix[dec] = dict(Counter(r["google_bucket"] for r in rows if r["anncsu_decision"] == dec))

    # Proxy agreement: among LOCAL_ACCEPT, how often Google had already accepted PVP,
    # vs had corrected/suspected it. This is diagnostic, not a truth metric.
    la = [r for r in rows if r["anncsu_decision"] == "LOCAL_ACCEPT"]
    accepted_by_google = sum(r["google_bucket"] in {"PVP_OK_CIVICO", "PVP_OK_STRADA"} for r in la)
    challenged_by_google = sum(r["google_bucket"] in {"CORRECTED_FROM_PVP", "PVP_SUSPECT", "PVP_UNCERTAIN"} for r in la)

    # Google correction comparisons against ANNCSU, only where both distances exist.
    comp = Counter()
    for r in rows:
        if r["google_bucket"] != "CORRECTED_FROM_PVP":
            continue
        a, b = r.get("distance_pvp_m"), r.get("distance_google_corrected_m")
        if not isinstance(a, (int, float)) or not isinstance(b, (int, float)):
            continue
        if a <= 25 and b > 50: comp["PVP_VICINO_ANNCSU_GOOGLE_LONTANO"] += 1
        elif b <= 25 and a > 50: comp["GOOGLE_VICINO_ANNCSU_PVP_LONTANO"] += 1
        elif a <= 25 and b <= 25: comp["ENTRAMBI_ENTRO_25M"] += 1
        elif a > 100 and b > 100: comp["ENTRAMBI_OLTRE_100M"] += 1
        elif b + 25 < a: comp["GOOGLE_PIU_VICINO_ANNCSU"] += 1
        elif a + 25 < b: comp["PVP_PIU_VICINO_ANNCSU"] += 1

    return {
        "version": VERSION,
        "generated_at": datetime.now().astimezone().isoformat(),
        "records": total,
        "municipality_unresolved": unresolved_count,
        "by_google_bucket": dict(by_google),
        "by_local_signal": dict(by_signal),
        "by_anncsu_decision": dict(by_decision),
        "by_decision_reason": dict(by_reason),
        "by_ambiguity": dict(by_ambiguity),
        "by_municipality_resolution": dict(by_muni_res),
        "by_province_relation": dict(province_rel),
        "decision_x_google_bucket": matrix,
        "local_accept": {
            "n": len(la),
            "pct": round(100.0 * len(la) / total, 2) if total else 0.0,
            "google_also_accepted_pvp": accepted_by_google,
            "google_challenged_pvp": challenged_by_google,
        },
        "google_correction_comparison": dict(comp),
        "notes": [
            "Nessun dato Supabase viene modificato.",
            "Nessuna API Google viene chiamata.",
            "v3 usa localita/frazione come disambiguazione positiva e non seleziona una strada candidata in base alla vicinanza al PVP.",
            "LOCAL_ACCEPT e' deliberatamente conservativo: civico esatto <=25m oppure strada senza civico <=50m, solo se ANNCSU non e' ambiguo e la provincia non e' incoerente.",
            "LOCAL_REVIEW indica evidenza ufficiale utile ma non sufficiente per auto-accettare il marker.",
        ],
    }


def main_v3() -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="Benchmark ANNCSU v3 locality/ambiguity-aware")
    ap.add_argument("--tools-dir", type=Path, default=here)
    ap.add_argument("--strad", type=Path)
    ap.add_argument("--indir", type=Path)
    ap.add_argument("--log", type=Path)
    ap.add_argument("--istat-xlsx", type=Path)
    ap.add_argument("--top-streets", type=int, default=6)
    args = ap.parse_args()

    base = args.tools_dir.resolve()
    strad = args.strad or pick_existing(["STRAD_ITA_*.csv"], base)
    indir = args.indir or pick_existing(["INDIR_ITA_*.csv"], base)
    log = args.log or pick_existing(["google_geocode_log_v8.jsonl", "google_geocode_log_v8*.jsonl"], base)
    istat = args.istat_xlsx or (base / "Elenco-comuni-italiani.xlsx")
    for label, path in (("STRAD", strad), ("INDIR", indir), ("LOG", log)):
        if path is None or not Path(path).exists():
            raise SystemExit(f"File {label} non trovato in {base}")
    strad, indir, log, istat = Path(strad), Path(indir), Path(log), Path(istat)

    print("=" * 76)
    print(f"Aste Easy - ANNCSU benchmark v{VERSION}")
    print(f"tools : {base}")
    print(f"STRAD : {strad.name}")
    print(f"INDIR : {indir.name}")
    print(f"LOG   : {log.name}")
    print("READ-ONLY: nessun update Supabase, nessuna chiamata Google.")
    print("=" * 76)

    t0 = time.time()
    records = load_repair_log(log)
    download_if_missing(istat, ISTAT_XLSX_URL)
    comuni = load_istat_comuni(istat)
    resolved, muni_meta, unresolved = resolve_municipalities_v3(records, comuni)
    top = scan_strad(strad, records, resolved, max(1, args.top_streets))
    wanted = {prog for arr in top.values() for _, prog, _, _ in arr}
    print(f"[MATCH] progressivi ANNCSU candidati complessivi: {len(wanted):,}")
    accesses = scan_indir(indir, wanted)

    result_rows = []
    for r in records:
        pid = str(r.get("pvp_id"))
        result_rows.append(evaluate_record_v3(r, resolved.get(pid), muni_meta.get(pid, {}), top.get(pid, []), accesses))

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_csv = base / f"anncsu_benchmark_v3_{stamp}.csv"
    out_json = base / f"anncsu_benchmark_summary_v3_{stamp}.json"
    out_unresolved = base / f"anncsu_comuni_irrisolti_v3_{stamp}.csv"
    out_conflicts = base / f"anncsu_google_conflicts_v3_{stamp}.csv"
    out_review = base / f"anncsu_review_cases_v3_{stamp}.csv"

    fields = list(result_rows[0].keys())
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(result_rows)

    summary = build_summary_v3(result_rows, len(unresolved))
    summary["files"] = {"strad": strad.name, "indir": indir.name, "google_log": log.name, "istat": istat.name}
    summary["elapsed_seconds"] = round(time.time() - t0, 2)

    with out_unresolved.open("w", encoding="utf-8-sig", newline="") as f:
        uf = ["pvp_id", "comune", "sigla_provincia", "cap", "indirizzo"]
        w = csv.DictWriter(f, fieldnames=uf); w.writeheader()
        for r in unresolved: w.writerow({k: r.get(k, "") for k in uf})

    conflict_rows = []
    for r in result_rows:
        if r.get("google_bucket") != "CORRECTED_FROM_PVP":
            continue
        d_old, d_new = r.get("distance_pvp_m"), r.get("distance_google_corrected_m")
        if not isinstance(d_old, (int, float)) or not isinstance(d_new, (int, float)):
            continue
        sig = ""
        if d_old <= 25 and d_new > 50: sig = "PVP_VICINO_ANNCSU_GOOGLE_LONTANO"
        elif d_new <= 25 and d_old > 50: sig = "GOOGLE_VICINO_ANNCSU_PVP_LONTANO"
        elif d_old <= 25 and d_new <= 25: sig = "ENTRAMBI_ENTRO_25M"
        elif d_old > 100 and d_new > 100: sig = "ENTRAMBI_OLTRE_100M"
        elif d_new + 25 < d_old: sig = "GOOGLE_PIU_VICINO_ANNCSU"
        elif d_old + 25 < d_new: sig = "PVP_PIU_VICINO_ANNCSU"
        if sig:
            rr = dict(r); rr["comparison_signal"] = sig; conflict_rows.append(rr)
    if conflict_rows:
        with out_conflicts.open("w", encoding="utf-8-sig", newline="") as f:
            cf = list(conflict_rows[0].keys()); w = csv.DictWriter(f, fieldnames=cf); w.writeheader(); w.writerows(conflict_rows)
    else:
        out_conflicts.write_text("", encoding="utf-8")

    review_rows = [r for r in result_rows if r["anncsu_decision"] != "LOCAL_ACCEPT"]
    with out_review.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(review_rows)

    summary["google_anncsu_conflict_rows"] = len(conflict_rows)
    summary["review_or_google_rows"] = len(review_rows)
    summary["files"].update({"conflicts": out_conflicts.name, "review": out_review.name})
    out_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n=== RISULTATO v3 ===")
    print(f"CSV dettagli : {out_csv}")
    print(f"JSON summary : {out_json}")
    print(f"Comuni dubbi : {out_unresolved}")
    print(f"Conflitti G/A: {out_conflicts}")
    print(f"Review/fallback: {out_review}")
    print("\nDecisioni simulate:")
    for k, v in sorted(summary["by_anncsu_decision"].items(), key=lambda kv: (-kv[1], kv[0])):
        print(f"  {k:20s} {v:6d}")
    print(f"\nLOCAL_ACCEPT: {summary['local_accept']['n']}/{summary['records']} = {summary['local_accept']['pct']}%")
    print("Non modifica ancora il repair: questa v3 serve a validare la regola definitiva.")
    return 0




# ============================================================================
# v5 overrides: direct use of high-accuracy official ANNCSU coordinates
# ============================================================================
VERSION = "5.0"
HIGH_ACCURACY_METHODS = {"1", "3"}
DIRECT_MIN_DISTANCE_M = 50.0
DIRECT_MIN_STREET_SCORE = 0.90


def _as_int(value, default=0):
    try:
        return int(float(value))
    except Exception:
        return default


def _locality_guard_v5(row: dict) -> str:
    """Conservative locality guard for direct correction.

    NO_HINT: source did not specify locality/frazione.
    CONFIRMED: ANNCSU locality explicitly matches a source hint.
    HINT_IN_ODONIMO: source locality is embedded in official street/toponym.
    UNCONFIRMED: source has locality context but ANNCSU has no matching locality.
    MISMATCH: ANNCSU locality exists and conflicts with the source hint.
    """
    raw_hints = str(row.get("source_locality_hints") or "").strip()
    if not raw_hints:
        return "NO_HINT"
    hints = [norm_text(x) for x in raw_hints.split("|") if norm_text(x)]
    ann_loc = norm_text(row.get("anncsu_localita"))
    loc_sc = safe_float(row.get("locality_score"))
    if ann_loc:
        if loc_sc is not None and loc_sc >= 0.84:
            return "CONFIRMED"
        return "MISMATCH"

    # Some rural addresses encode the locality directly in the official odonimo
    # (e.g. VOCABOLO/LOCALITA' POGGIO PARRANO) while ANNCSU LOCALITA' is blank.
    odo_tokens = set(tokens(row.get("anncsu_odonimo"), drop_stop=True))
    for h in hints:
        ht = set(tokens(h, drop_stop=True))
        if not ht:
            continue
        overlap = len(ht & odo_tokens) / max(1, len(ht))
        if overlap >= 0.75:
            return "HINT_IN_ODONIMO"
    return "UNCONFIRMED"


def _benchmark_signal_v5(row: dict) -> str:
    bucket = str(row.get("google_bucket") or "")
    dnew = safe_float(row.get("distance_google_corrected_m"))
    if bucket == "CORRECTED_FROM_PVP":
        if dnew is None:
            return "GOOGLE_CORRECTED_NO_COMPARISON"
        if dnew <= 10:
            return "GOOGLE_CORRECTION_WITHIN_10M_ANNCSU"
        if dnew <= 25:
            return "GOOGLE_CORRECTION_WITHIN_25M_ANNCSU"
        if dnew <= 50:
            return "GOOGLE_CORRECTION_WITHIN_50M_ANNCSU"
        if dnew <= 100:
            return "GOOGLE_CORRECTION_50_100M_FROM_ANNCSU"
        return "GOOGLE_CORRECTION_GT100M_FROM_ANNCSU"
    if bucket in {"PVP_OK_CIVICO", "PVP_OK_STRADA"}:
        return "GOOGLE_ACCEPTED_OLD_PVP"
    if bucket == "PVP_SUSPECT":
        return "GOOGLE_LEFT_SUSPECT"
    if bucket == "PVP_UNCERTAIN":
        return "GOOGLE_LEFT_UNCERTAIN"
    return bucket or "OTHER"


def evaluate_record_v5(r: dict, code: Optional[str], municipality_meta: dict, street_cands, accesses) -> dict:
    # Start from the validated v3 semantic/ambiguity logic. No neighbor logic.
    out = evaluate_record_v3(r, code, municipality_meta, street_cands, accesses)
    out["v3_anncsu_decision"] = out.get("anncsu_decision", "")
    out["v3_decision_reason"] = out.get("decision_reason", "")
    out["direct_locality_guard"] = _locality_guard_v5(out)
    out["direct_candidate"] = "NO"
    out["direct_candidate_reason"] = ""
    out["direct_anncsu_lat"] = ""
    out["direct_anncsu_lon"] = ""
    out["direct_anncsu_method"] = ""
    out["direct_distance_pvp_m"] = ""
    out["direct_distance_google_m"] = ""
    out["benchmark_signal"] = ""

    # Preserve a v3 LOCAL_ACCEPT: no correction is needed.
    if out.get("anncsu_decision") == "LOCAL_ACCEPT":
        out["benchmark_signal"] = _benchmark_signal_v5(out)
        return out

    sc = safe_float(out.get("street_score"))
    d_old = safe_float(out.get("distance_pvp_m"))
    lat = safe_float(out.get("anncsu_lat"))
    lon = safe_float(out.get("anncsu_lon"))
    method = str(out.get("anncsu_metodo") or "").strip()
    if method.endswith(".0"):
        method = method[:-2]
    access_count = _as_int(out.get("matching_access_candidates"), 0)
    locality_guard = out["direct_locality_guard"]

    structural_ok = (
        out.get("province_relation") == "MATCH"
        and out.get("ambiguity") == "NO"
        and out.get("civic_relation") == "exact"
        and sc is not None and sc >= DIRECT_MIN_STREET_SCORE
        and access_count == 1
        and method in HIGH_ACCURACY_METHODS
        and lat is not None and lon is not None
        and d_old is not None
    )

    if not structural_ok:
        # Add a useful diagnostic without changing the v3 fallback semantics.
        reasons = []
        if out.get("province_relation") != "MATCH": reasons.append("province_not_match")
        if out.get("ambiguity") != "NO": reasons.append("anncsu_ambiguous")
        if out.get("civic_relation") != "exact": reasons.append("civic_not_exact")
        if sc is None or sc < DIRECT_MIN_STREET_SCORE: reasons.append("street_score_below_0_90")
        if access_count != 1: reasons.append("exact_access_not_unique")
        if method not in HIGH_ACCURACY_METHODS: reasons.append("method_not_1_3")
        if lat is None or lon is None: reasons.append("missing_anncsu_coordinates")
        if d_old is None: reasons.append("missing_pvp_distance")
        out["direct_candidate_reason"] = "|".join(reasons) or "not_structurally_eligible"
        out["benchmark_signal"] = _benchmark_signal_v5(out)
        return out

    out["direct_anncsu_lat"] = lat
    out["direct_anncsu_lon"] = lon
    out["direct_anncsu_method"] = method
    out["direct_distance_pvp_m"] = round(d_old, 2)
    d_google = safe_float(out.get("distance_google_corrected_m"))
    if d_google is not None:
        out["direct_distance_google_m"] = round(d_google, 2)

    # Explicit locality mismatch never permits direct correction. If the source
    # specified a locality that ANNCSU cannot independently confirm, keep review.
    if locality_guard == "MISMATCH":
        out["anncsu_decision"] = "GOOGLE_NEEDED"
        out["decision_reason"] = "high_accuracy_civic_but_locality_mismatch"
        out["direct_candidate_reason"] = "locality_mismatch"
    elif locality_guard == "UNCONFIRMED":
        out["anncsu_decision"] = "LOCAL_REVIEW"
        out["decision_reason"] = "high_accuracy_civic_but_locality_unconfirmed"
        out["direct_candidate_reason"] = "locality_unconfirmed"
    elif d_old <= 25:
        # Normally already caught by v3 LOCAL_ACCEPT; retained for completeness.
        out["anncsu_decision"] = "LOCAL_ACCEPT"
        out["decision_reason"] = "high_accuracy_exact_civic_within_25m"
        out["direct_candidate_reason"] = "already_close"
    elif d_old <= DIRECT_MIN_DISTANCE_M:
        out["anncsu_decision"] = "LOCAL_REVIEW"
        out["decision_reason"] = "high_accuracy_exact_civic_25_50m_review"
        out["direct_candidate_reason"] = "movement_too_small_for_direct_correction"
    else:
        out["anncsu_decision"] = "LOCAL_CORRECT_ANNCSU"
        out["decision_reason"] = "unique_high_accuracy_anncsu_civic_over_50m_from_pvp"
        out["direct_candidate"] = "YES"
        out["direct_candidate_reason"] = "strong_direct_coordinate_candidate"

    out["benchmark_signal"] = _benchmark_signal_v5(out)
    return out


def _median(nums):
    vals = sorted(float(x) for x in nums if isinstance(x, (int, float)) and math.isfinite(float(x)))
    if not vals:
        return None
    n = len(vals)
    return vals[n // 2] if n % 2 else (vals[n // 2 - 1] + vals[n // 2]) / 2.0


def build_summary_v5(rows: List[dict], unresolved_count: int) -> dict:
    total = len(rows)
    by_google = Counter(r.get("google_bucket") for r in rows)
    by_signal = Counter(r.get("local_signal") for r in rows)
    by_decision = Counter(r.get("anncsu_decision") for r in rows)
    by_reason = Counter(r.get("decision_reason") for r in rows)
    by_ambiguity = Counter(r.get("ambiguity") or "N/A" for r in rows)
    by_muni_res = Counter(r.get("municipality_resolution") for r in rows)
    province_rel = Counter(r.get("province_relation") for r in rows)
    locality_guard = Counter(r.get("direct_locality_guard") or "N/A" for r in rows)

    matrix = {}
    for dec in ["LOCAL_ACCEPT", "LOCAL_CORRECT_ANNCSU", "LOCAL_REVIEW", "GOOGLE_NEEDED"]:
        matrix[dec] = dict(Counter(r.get("google_bucket") for r in rows if r.get("anncsu_decision") == dec))

    accepts = [r for r in rows if r.get("anncsu_decision") == "LOCAL_ACCEPT"]
    directs = [r for r in rows if r.get("anncsu_decision") == "LOCAL_CORRECT_ANNCSU"]
    zero_google = accepts + directs
    direct_google_corrected = [r for r in directs if r.get("google_bucket") == "CORRECTED_FROM_PVP"]
    dg = [safe_float(r.get("distance_google_corrected_m")) for r in direct_google_corrected]
    dg = [x for x in dg if x is not None]
    dp = [safe_float(r.get("distance_pvp_m")) for r in directs]
    dp = [x for x in dp if x is not None]

    direct_benchmark_signals = Counter(r.get("benchmark_signal") or "N/A" for r in directs)
    direct_methods = Counter(str(r.get("direct_anncsu_method") or "") for r in directs)

    return {
        "version": VERSION,
        "generated_at": datetime.now().astimezone().isoformat(),
        "records": total,
        "municipality_unresolved": unresolved_count,
        "by_google_bucket": dict(by_google),
        "by_local_signal": dict(by_signal),
        "by_anncsu_decision": dict(by_decision),
        "by_decision_reason": dict(by_reason),
        "by_ambiguity": dict(by_ambiguity),
        "by_municipality_resolution": dict(by_muni_res),
        "by_province_relation": dict(province_rel),
        "by_direct_locality_guard": dict(locality_guard),
        "decision_x_google_bucket": matrix,
        "zero_google_simulation": {
            "local_accept_n": len(accepts),
            "local_correct_anncsu_n": len(directs),
            "total_n": len(zero_google),
            "pct": round(100.0 * len(zero_google) / total, 2) if total else 0.0,
        },
        "direct_correction_benchmark": {
            "n": len(directs),
            "by_google_bucket": dict(Counter(r.get("google_bucket") for r in directs)),
            "by_anncsu_method": dict(direct_methods),
            "by_benchmark_signal": dict(direct_benchmark_signals),
            "median_pvp_to_anncsu_m": None if _median(dp) is None else round(_median(dp), 2),
            "google_corrected_comparable_n": len(dg),
            "median_google_to_anncsu_m": None if _median(dg) is None else round(_median(dg), 2),
            "google_to_anncsu_le_10m": sum(x <= 10 for x in dg),
            "google_to_anncsu_le_25m": sum(x <= 25 for x in dg),
            "google_to_anncsu_le_50m": sum(x <= 50 for x in dg),
            "google_to_anncsu_le_100m": sum(x <= 100 for x in dg),
            "google_to_anncsu_gt_100m": sum(x > 100 for x in dg),
        },
        "rules": {
            "local_accept": "v3 unchanged: exact civic <=25m or no-civic street <=50m, unambiguous/province coherent",
            "local_correct_anncsu": "exact civic, street_score>=0.90, exactly one ANNCSU exact access, METODO 1/3, coordinates present, no ambiguity, province match, locality absent/confirmed, PVP >50m",
            "direct_review_band": "25m < PVP distance <=50m remains LOCAL_REVIEW",
        },
        "notes": [
            "Nessun dato Supabase viene modificato.",
            "Nessuna API Google viene chiamata.",
            "v5 parte dalla v3; la neighbor evidence v4 NON viene usata per auto-accettare.",
            "LOCAL_CORRECT_ANNCSU e' una simulazione di benchmark, non modifica ancora il DB.",
            "La coordinata diretta e' ammessa solo con un singolo accesso esatto ANNCSU, evitando di scegliere tra coordinate multiple in base alla vicinanza al PVP.",
            "METODO 1/3 sono trattati come high-accuracy secondo i metadata ANNCSU forniti per il benchmark.",
        ],
    }


def _write_rows(path: Path, rows: List[dict], fields: List[str]):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        if rows:
            w.writerows(rows)


def main_v5() -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="Benchmark ANNCSU v5 direct high-accuracy coordinates")
    ap.add_argument("--tools-dir", type=Path, default=here)
    ap.add_argument("--strad", type=Path)
    ap.add_argument("--indir", type=Path)
    ap.add_argument("--log", type=Path)
    ap.add_argument("--istat-xlsx", type=Path)
    ap.add_argument("--top-streets", type=int, default=6)
    args = ap.parse_args()

    base = args.tools_dir.resolve()
    strad = args.strad or pick_existing(["STRAD_ITA_*.csv"], base)
    indir = args.indir or pick_existing(["INDIR_ITA_*.csv"], base)
    log = args.log or pick_existing(["google_geocode_log_v8.jsonl", "google_geocode_log_v8*.jsonl"], base)
    istat = args.istat_xlsx or (base / "Elenco-comuni-italiani.xlsx")
    for label, path in (("STRAD", strad), ("INDIR", indir), ("LOG", log)):
        if path is None or not Path(path).exists():
            raise SystemExit(f"File {label} non trovato in {base}")
    strad, indir, log, istat = Path(strad), Path(indir), Path(log), Path(istat)

    print("=" * 78)
    print(f"Aste Easy - ANNCSU benchmark v{VERSION} direct-coordinate")
    print(f"tools : {base}")
    print(f"STRAD : {strad.name}")
    print(f"INDIR : {indir.name}")
    print(f"LOG   : {log.name}")
    print("READ-ONLY: nessun update Supabase, nessuna chiamata Google.")
    print("=" * 78)

    t0 = time.time()
    records = load_repair_log(log)
    download_if_missing(istat, ISTAT_XLSX_URL)
    comuni = load_istat_comuni(istat)
    resolved, muni_meta, unresolved = resolve_municipalities_v3(records, comuni)
    top = scan_strad(strad, records, resolved, max(1, args.top_streets))
    wanted = {prog for arr in top.values() for _, prog, _, _ in arr}
    print(f"[MATCH] progressivi ANNCSU candidati complessivi: {len(wanted):,}")
    accesses = scan_indir(indir, wanted)

    result_rows = []
    for r in records:
        pid = str(r.get("pvp_id"))
        result_rows.append(evaluate_record_v5(r, resolved.get(pid), muni_meta.get(pid, {}), top.get(pid, []), accesses))

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_csv = base / f"anncsu_benchmark_v5_{stamp}.csv"
    out_json = base / f"anncsu_benchmark_summary_v5_{stamp}.json"
    out_unresolved = base / f"anncsu_comuni_irrisolti_v5_{stamp}.csv"
    out_conflicts = base / f"anncsu_google_conflicts_v5_{stamp}.csv"
    out_review = base / f"anncsu_review_cases_v5_{stamp}.csv"
    out_direct = base / f"anncsu_direct_corrections_v5_{stamp}.csv"
    out_outliers = base / f"anncsu_direct_outliers_v5_{stamp}.csv"

    fields = list(result_rows[0].keys())
    _write_rows(out_csv, result_rows, fields)

    uf = ["pvp_id", "comune", "sigla_provincia", "cap", "indirizzo"]
    unresolved_rows = [{k: r.get(k, "") for k in uf} for r in unresolved]
    _write_rows(out_unresolved, unresolved_rows, uf)

    # Keep the v3 Google/ANNCSU conflict report for continuity.
    conflict_rows = []
    for r in result_rows:
        if r.get("google_bucket") != "CORRECTED_FROM_PVP":
            continue
        d_old = safe_float(r.get("distance_pvp_m")); d_new = safe_float(r.get("distance_google_corrected_m"))
        if d_old is None or d_new is None:
            continue
        sig = ""
        if d_old <= 25 and d_new > 50: sig = "PVP_VICINO_ANNCSU_GOOGLE_LONTANO"
        elif d_new <= 25 and d_old > 50: sig = "GOOGLE_VICINO_ANNCSU_PVP_LONTANO"
        elif d_old <= 25 and d_new <= 25: sig = "ENTRAMBI_ENTRO_25M"
        elif d_old > 100 and d_new > 100: sig = "ENTRAMBI_OLTRE_100M"
        elif d_new + 25 < d_old: sig = "GOOGLE_PIU_VICINO_ANNCSU"
        elif d_old + 25 < d_new: sig = "PVP_PIU_VICINO_ANNCSU"
        if sig:
            rr = dict(r); rr["comparison_signal"] = sig; conflict_rows.append(rr)
    conflict_fields = fields + (["comparison_signal"] if "comparison_signal" not in fields else [])
    _write_rows(out_conflicts, conflict_rows, conflict_fields)

    direct_rows = [r for r in result_rows if r.get("anncsu_decision") == "LOCAL_CORRECT_ANNCSU"]
    _write_rows(out_direct, direct_rows, fields)

    # Outliers/disagreements worth manual inspection before any production write.
    outlier_rows = []
    for r in direct_rows:
        bucket = r.get("google_bucket")
        dnew = safe_float(r.get("distance_google_corrected_m"))
        reason = ""
        if bucket in {"PVP_OK_CIVICO", "PVP_OK_STRADA"}:
            reason = "google_had_accepted_old_pvp"
        elif bucket == "CORRECTED_FROM_PVP" and dnew is not None and dnew > 50:
            reason = "google_correction_over_50m_from_anncsu"
        elif bucket in {"PVP_SUSPECT", "PVP_UNCERTAIN"}:
            reason = "google_left_unresolved"
        if reason:
            rr = dict(r); rr["outlier_reason"] = reason; outlier_rows.append(rr)
    outlier_fields = fields + (["outlier_reason"] if "outlier_reason" not in fields else [])
    _write_rows(out_outliers, outlier_rows, outlier_fields)

    review_rows = [r for r in result_rows if r.get("anncsu_decision") in {"LOCAL_REVIEW", "GOOGLE_NEEDED"}]
    _write_rows(out_review, review_rows, fields)

    summary = build_summary_v5(result_rows, len(unresolved))
    summary["files"] = {
        "strad": strad.name, "indir": indir.name, "google_log": log.name, "istat": istat.name,
        "conflicts": out_conflicts.name, "review": out_review.name,
        "direct_corrections": out_direct.name, "direct_outliers": out_outliers.name,
    }
    summary["elapsed_seconds"] = round(time.time() - t0, 2)
    summary["google_anncsu_conflict_rows"] = len(conflict_rows)
    summary["direct_outlier_rows"] = len(outlier_rows)
    summary["review_or_google_rows"] = len(review_rows)
    out_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n=== RISULTATO v5 ===")
    print(f"CSV dettagli      : {out_csv}")
    print(f"JSON summary      : {out_json}")
    print(f"Comuni dubbi      : {out_unresolved}")
    print(f"Conflitti G/A     : {out_conflicts}")
    print(f"Review/fallback   : {out_review}")
    print(f"Correzioni ANNCSU : {out_direct}")
    print(f"Outlier diretti   : {out_outliers}")
    print("\nDecisioni simulate:")
    for k, v in sorted(summary["by_anncsu_decision"].items(), key=lambda kv: (-kv[1], kv[0])):
        print(f"  {k:24s} {v:6d}")
    zg = summary["zero_google_simulation"]
    print(f"\nZERO-GOOGLE simulato: {zg['total_n']}/{summary['records']} = {zg['pct']}%")
    dc = summary["direct_correction_benchmark"]
    print(f"LOCAL_CORRECT_ANNCSU: {dc['n']} (Google-corrected comparabili: {dc['google_corrected_comparable_n']})")
    print("Ancora benchmark: NON usare questi output per aggiornare Supabase.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main_v5())
