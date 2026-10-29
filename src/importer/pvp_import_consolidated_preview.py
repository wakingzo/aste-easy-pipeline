import argparse
import csv
import hashlib
import math
import unicodedata
from collections import Counter
from urllib.parse import urlencode
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen

BASE_DIR = Path(__file__).resolve().parent
SRC_ROOT = BASE_DIR.parent
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))
from aste_paths import RAW_DIR, REFERENCE_DIR, REPORTS_DIR, ensure_runtime_dirs
from project_env import load_project_env

ensure_runtime_dirs()
IN_RAW = RAW_DIR / "listings_api_raw.jsonl"
IN_ALLEGATI = RAW_DIR / "allegati_links.jsonl"
IN_SOGGETTI = RAW_DIR / "listing_soggetti.jsonl"
COMUNI_CSV = REFERENCE_DIR / "comuni_cap.csv"

OUT_JSONL = RAW_DIR / "hybrid_api_rows_full.jsonl"

SUPABASE_URL = os.environ.get("SUPABASE_URL", "").strip()
SUPABASE_SERVICE_ROLE = os.environ.get("SUPABASE_SERVICE_ROLE", "").strip()
SUPABASE_ENABLED = bool(SUPABASE_URL and SUPABASE_SERVICE_ROLE)

BATCH_SIZE_LISTINGS = int(os.environ.get("BATCH_SIZE_LISTINGS", "200"))
BATCH_SIZE_ALLEGATI = int(os.environ.get("BATCH_SIZE_ALLEGATI", "500"))
BATCH_SIZE_DELETE = int(os.environ.get("BATCH_SIZE_DELETE", "200"))

DISP_MAP = {
    "occup": "Occupato",
    "occupato": "Occupato",
    "liber": "Libero",
    "libero": "Libero",
    "occst": "Occupato senza titolo",
    "occupato_senza_titolo": "Occupato senza titolo",
    "non_specificato": "",
    "non_specificata": "",
    "": "",
}

LISTINGS_FIELDS = [
    "pvp_id", "link_pvp", "data_pubblicazione", "data_pubblicazione_date", "data_asta",
    "regione", "comune", "provincia", "sigla_provincia", "cap", "indirizzo", "maps_link",
    "latitudine", "longitudine", "tipologia", "tribunale", "tipo_procedura", "numero_procedura",
    "anno_procedura", "categoria_procedura", "categoria_procedura_reason", "lotto", "descrizione",
    "categoria_immobile", "categoria_immobile_reason", "prezzo_base", "offerta_minima",
    "rialzo_minimo", "disponibilita", "vani", "piano", "superficie_mq", "esperimento_vendita",
    "note", "status"
]


def load_jsonl(path: Path):
    out = []
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8", errors="replace") as f:
        for line in f:
            line = line.strip()
            if line:
                out.append(json.loads(line))
    return out


def clean(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\u00A0", " ")
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def first_non_empty(*vals):
    for v in vals:
        if v is None:
            continue
        if isinstance(v, (list, tuple)):
            for item in v:
                s = clean(item)
                if s and s.upper() not in {"NULL", "NONE", "NON_SPECIFICATO", "NON_SPECIFICATA", "-"}:
                    return s
            continue
        s = clean(v)
        if s and s.upper() not in {"NULL", "NONE", "NON_SPECIFICATO", "NON_SPECIFICATA", "-"}:
            return s
    return ""


def titleish(s: str) -> str:
    s = clean(s).replace("_", " ").replace("-", " ")
    return " ".join(w.capitalize() if w.isalpha() else w for w in s.split())


def euro_it(v) -> str:
    if v is None or v == "":
        return ""
    try:
        x = float(v)
    except Exception:
        return clean(v)
    s = f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s + " €"


def date_it(v: str) -> str:
    s = clean(v).replace("Z", "")
    if not s:
        return ""
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y", "%d/%m/%Y %H:%M"):
        try:
            dt = datetime.strptime(s[:19] if "T" in s and len(s) >= 19 else s, fmt)
            return dt.strftime("%d/%m/%Y")
        except Exception:
            pass
    m = re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b", s)
    return f"{m.group(3)}/{m.group(2)}/{m.group(1)}" if m else ""


def iso_date_from_it(v: str) -> str:
    s = clean(v)
    if not s:
        return ""
    try:
        return datetime.strptime(s, "%d/%m/%Y").strftime("%Y-%m-%d")
    except Exception:
        return ""


def normalize_disponibilita(v) -> str:
    if v is None:
        return ""
    vals = v if isinstance(v, list) else [v]
    out = []
    for x in vals:
        s = clean(x)
        if not s:
            continue
        mapped = DISP_MAP.get(s.lower().replace(" ", "_"), titleish(s))
        if mapped and mapped not in out:
            out.append(mapped)
    return " | ".join(out)


def prettify_api_category(v) -> str:
    s = clean(v)
    return titleish(s) if s else ""


def build_maps_link(lat, lon, indirizzo="", comune="", provincia=""):
    lat_s = clean(lat)
    lon_s = clean(lon)
    if lat_s and lon_s:
        return f"https://www.google.com/maps?q={lat_s},{lon_s}"
    q = " ".join(x for x in [clean(indirizzo), clean(comune), clean(provincia)] if clean(x))
    return f"https://www.google.com/maps?q={quote(q)}" if q else ""


def clean_esperimento(v):
    s = clean(v).replace(" ?", "")
    if "/" in s and re.fullmatch(r"[IVXLC]+/[IVXLC]+", s):
        return s.split("/")[0]
    return s


def _roman(n: int) -> str:
    return {1:"I",2:"II",3:"III",4:"IV",5:"V",6:"VI",7:"VII",8:"VIII",9:"IX",10:"X"}.get(n, "")


def extract_esperimento_from_filenames(files):
    if not files:
        return ""
    for f in files:
        up = f.upper()
        for rn in [" X ", " IX ", " VIII ", " VII ", " VI ", " V ", " IV ", " III ", " II ", " I "]:
            if rn.strip() in up:
                return rn.strip()
        m = re.search(r"\b(10|[1-9])\s*[°ºª\^]?", f)
        if m:
            return _roman(int(m.group(1)))
    return ""


def load_comuni_lookup(csv_path: Path):
    comuni_map = {}
    cap_map = {}
    try:
        f = open(csv_path, "r", encoding="utf-8-sig", errors="replace", newline="")
    except Exception:
        f = open(csv_path, "r", encoding="latin-1", errors="replace", newline="")
    with f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            comune = clean(row.get("comune") or row.get("Comune") or row.get("ï»¿comune") or "")
            if not comune:
                continue
            info = {
                "regione": clean(row.get("regione") or row.get("Regione") or ""),
                "provincia": clean(row.get("provincia") or row.get("Provincia") or ""),
                "sigla_provincia": clean(row.get("sigla_provincia") or row.get("sigla") or row.get("Sigla") or ""),
                "comune": comune,
                "cap": re.sub(r"\D", "", clean(row.get("cap") or row.get("CAP") or "")).zfill(5),
            }
            comuni_map.setdefault(comune.lower(), []).append(info)
            if info["cap"]:
                cap_map.setdefault(info["cap"], []).append(info)
    return comuni_map, cap_map


COMUNI_MAP, CAP_MAP = load_comuni_lookup(COMUNI_CSV)


def match_comune_info(indirizzo: str):
    s = clean(indirizzo).lower()
    if not s:
        return {}
    m = re.search(r"\b(\d{5})\b", s)
    if m and m.group(1) in CAP_MAP:
        return CAP_MAP[m.group(1)][0]
    for comune, infos in COMUNI_MAP.items():
        if comune and comune in s:
            return infos[0]
    return {}




def canonical_text_key(s: str) -> str:
    s = clean(s).lower()
    s = s.replace("’", "'")
    s = s.replace("`", "'")
    s = s.replace("ü", "u").replace("ö", "o").replace("ä", "a")
    s = re.sub(r"\s+", " ", s)
    return s.strip()


REGIONE_ALIASES = {
    "trentino-alto adige": "Trentino-Alto Adige/Südtirol",
    "trentino alto adige": "Trentino-Alto Adige/Südtirol",
    "trentino-alto-adige": "Trentino-Alto Adige/Südtirol",
    "trentino-alto adige/sudtirol": "Trentino-Alto Adige/Südtirol",
    "trentino alto adige/sudtirol": "Trentino-Alto Adige/Südtirol",
    "trentino-alto adige/sud tirol": "Trentino-Alto Adige/Südtirol",
    "trentino alto adige/sud tirol": "Trentino-Alto Adige/Südtirol",
    "trentino-alto adige/südtirol": "Trentino-Alto Adige/Südtirol",
    "trentino alto adige/südtirol": "Trentino-Alto Adige/Südtirol",
}

PROVINCIA_ALIASES = {
    "reggio di calabria": "Reggio Calabria",
}


def canonicalize_geo_labels(regione="", provincia="", sigla=""):
    regione_clean = clean(regione)
    provincia_clean = clean(provincia)
    sigla_clean = clean(sigla).upper()

    reg_key = canonical_text_key(regione_clean)
    prov_key = canonical_text_key(provincia_clean)

    regione_final = REGIONE_ALIASES.get(reg_key, regione_clean)
    provincia_final = PROVINCIA_ALIASES.get(prov_key, provincia_clean)

    if provincia_final == "Reggio Calabria" and not sigla_clean:
        sigla_clean = "RC"

    return {
        "regione": regione_final,
        "provincia": provincia_final,
        "sigla_provincia": sigla_clean,
    }


def normalize_geo_with_csv(comune="", provincia="", sigla="", regione="", cap=""):
    comune = clean(comune)
    provincia = clean(provincia)
    sigla = clean(sigla)
    regione = clean(regione)
    cap = re.sub(r"\D", "", clean(cap))
    if cap:
        cap = cap.zfill(5)
    if not comune:
        return {
            "comune": comune,
            "provincia": provincia,
            "sigla_provincia": sigla,
            "regione": regione,
            "cap": cap,
        }
    candidates = COMUNI_MAP.get(comune.lower(), [])
    chosen = None
    for info in candidates:
        if provincia and clean(info.get("provincia")).lower() == provincia.lower():
            chosen = info
            break
        if sigla and clean(info.get("sigla_provincia")).lower() == sigla.lower():
            chosen = info
            break
        if regione and clean(info.get("regione")).lower() == regione.lower():
            chosen = info
            break
        if cap and clean(info.get("cap")) == cap:
            chosen = info
            break
    if not chosen and len(candidates) == 1:
        chosen = candidates[0]
    if not chosen:
        return {
            "comune": comune,
            "provincia": provincia,
            "sigla_provincia": sigla,
            "regione": regione,
            "cap": cap,
        }
    return {
        "comune": comune or clean(chosen.get("comune")),
        "provincia": clean(chosen.get("provincia")) or provincia,
        "sigla_provincia": clean(chosen.get("sigla_provincia")) or sigla,
        "regione": clean(chosen.get("regione")) or regione,
        "cap": clean(chosen.get("cap")) or cap,
    }


def clean_via_structured(via, comune="", provincia="", sigla="", cap="") -> str:
    s = clean(via)
    if not s:
        return ""
    s = s.replace("lLoc.", "Loc.").replace("lLocalità", "Località")
    s = re.sub(r"\blloc\.", "loc.", s, flags=re.I)
    s = re.sub(r"\bllocalit[àa]\b", "località", s, flags=re.I)
    s = re.sub(r",?\s*Italia\s*$", "", s, flags=re.I)
    s = re.sub(r"\bProvincia\s+(?:di|della|del)\b.*$", "", s, flags=re.I)
    s = re.sub(r"\b(Nord-Est Sardegna|Sud Sardegna)\b.*$", "", s, flags=re.I)
    comune_esc = re.escape(clean(comune))
    if comune:
        s = re.sub(
            rf"^\s*{comune_esc}\s+(?=(?:l?loc\.?|l?località|contrada|c\.da|via|viale|piazza|corso|vicolo|strada)\b)",
            "",
            s,
            flags=re.I,
        )
    prov_esc = re.escape(clean(provincia))
    sigla_esc = re.escape(clean(sigla))
    cap_esc = re.escape(clean(cap))
    patterns = []
    if cap and comune and sigla:
        patterns.append(rf",?\s*{cap_esc}\s+{comune_esc}\s+{sigla_esc}\s*$")
    if cap and comune and provincia:
        patterns.append(rf",?\s*{cap_esc}\s+{comune_esc}\s+{prov_esc}\s*$")
    if comune and sigla:
        patterns.append(rf",?\s*{comune_esc}\s+{sigla_esc}\s*$")
    if comune and provincia:
        patterns.append(rf",?\s*{comune_esc}\s+{prov_esc}\s*$")
    if cap and comune:
        patterns.append(rf",?\s*{cap_esc}\s+{comune_esc}\s*$")
    changed = True
    while changed:
        changed = False
        for pat in patterns:
            ns = re.sub(pat, "", s, flags=re.I).strip(" ,")
            if ns != s:
                s = ns
                changed = True
    s = re.sub(r"\s*,\s*,+", ", ", s)
    s = re.sub(r"\s+", " ", s).strip(" ,")
    parts = [p.strip(" ,") for p in re.split(r"\s*,\s*", s) if p.strip(" ,")]
    dedup = []
    seen_norm = set()
    for p in parts:
        norm = re.sub(r"\W+", "", p).lower()
        if not norm or norm in seen_norm:
            continue
        seen_norm.add(norm)
        dedup.append(p)
    return ", ".join(dedup)


def compose_indirizzo(via="", cap="", comune="", sigla="") -> str:
    line1 = clean(via)
    line2 = " ".join(x for x in [clean(cap), clean(comune), clean(sigla)] if clean(x)).strip()
    return ", ".join(x for x in [line1, line2] if x)


def build_address_from_bene(bene: dict) -> str:
    addr = bene.get("indirizzo") or {}
    comune = clean(addr.get("descComune") or addr.get("citta") or addr.get("comune") or "")
    provincia = clean(addr.get("descProvincia") or addr.get("provincia") or "")
    sigla = clean(addr.get("siglaProvincia") or addr.get("sigla_provincia") or "")
    cap = re.sub(r"\D", "", clean(addr.get("cap") or "")).zfill(5) if clean(addr.get("cap") or "") else ""
    via_raw = clean(addr.get("via") or addr.get("descVia") or "")
    nc = clean(addr.get("numeroCivico") or "")
    via = clean_via_structured(via_raw, comune, provincia, sigla, cap)
    if nc and nc not in via:
        via = " ".join(x for x in [via, nc] if x).strip()
    return compose_indirizzo(via, cap, comune, sigla or provincia)


def extract_geo_api(detail_body: dict, search_item: dict, bene_list: list):
    out = {"comune":"", "provincia":"", "sigla_provincia":"", "regione":"", "cap":"", "lat":"", "lon":""}
    addr_sources = []
    for b in bene_list:
        addr_sources.append(b.get("indirizzo") or {})
    addr_sources.append(search_item.get("indirizzo") or {})
    addr_sources.append(detail_body.get("indirizzo") or {})
    for addr in addr_sources:
        if not isinstance(addr, dict):
            continue
        coords = addr.get("coordinate") or {}
        out["comune"] = out["comune"] or first_non_empty(addr.get("descComune"), addr.get("citta"), addr.get("comune"))
        out["provincia"] = out["provincia"] or first_non_empty(addr.get("descProvincia"), addr.get("provincia"))
        out["sigla_provincia"] = out["sigla_provincia"] or first_non_empty(addr.get("siglaProvincia"), addr.get("sigla_provincia"))
        out["regione"] = out["regione"] or first_non_empty(addr.get("descRegione"), addr.get("regione"), addr.get("codRegione"))
        cap = re.sub(r"\D", "", clean(addr.get("cap") or "")).zfill(5) if clean(addr.get("cap") or "") else ""
        out["cap"] = out["cap"] or cap
        out["lat"] = out["lat"] or first_non_empty(addr.get("latitudine"), addr.get("lat"), coords.get("latitudine"), coords.get("lat"))
        out["lon"] = out["lon"] or first_non_empty(addr.get("longitudine"), addr.get("lon"), addr.get("lng"), coords.get("longitudine"), coords.get("lon"), coords.get("lng"))
    return out


def pick_categoria_immobile(detail_body: dict, search_item: dict, fallback=""):
    lotto = detail_body.get("lotto") or {}
    for cand in [
        detail_body.get("categoriaLotto"), detail_body.get("descTipoCategLotto"), detail_body.get("codTipoCategLotto"),
        lotto.get("categoriaLotto"), lotto.get("descTipoCategLotto"), lotto.get("codTipoCategLotto"),
        search_item.get("ncategoriaLotto"), search_item.get("categoriaLotto"),
    ]:
        s = prettify_api_category(cand)
        if s:
            return s, "api_categoria_lotto"
    return fallback, "fallback"


def pick_tipologia(detail_body: dict, search_item: dict, fallback=""):
    lotto = detail_body.get("lotto") or {}
    beni = (detail_body.get("beni") or []) + (lotto.get("beni") or [])
    for b in beni:
        for k in ("categoriaBene", "descCategoriaBene", "tipoBene", "descTipologiaBene", "tipologia", "codTipologiaBene"):
            s = prettify_api_category(b.get(k))
            if s:
                return s
    for k in ("categoriaBene", "descCategoriaBene", "tipoBene", "descTipologiaBene", "tipologia"):
        s = prettify_api_category(search_item.get(k))
        if s:
            return s
    return fallback


def pick_tipo_procedura(detail_body: dict, search_item: dict, fallback=""):
    procedura = detail_body.get("procedura") or {}
    return first_non_empty(procedura.get("tipoRito"), procedura.get("descTipoRito"), detail_body.get("tipoRito"), detail_body.get("descTipoRito"), search_item.get("tipoRito"), search_item.get("descTipoRito"), fallback)


def pick_categoria_procedura(detail_body: dict, search_item: dict, fallback=""):
    procedura = detail_body.get("procedura") or {}
    s = first_non_empty(procedura.get("tipoRegistro"), procedura.get("descTipoRegistro"), detail_body.get("tipoRegistro"), detail_body.get("descTipoRegistro"), search_item.get("tipoRegistro"), search_item.get("descTipoRegistro"), fallback)
    return s, ("api_tipo_registro" if s and s != fallback else "fallback")


def pick_data_pubblicazione(detail_body: dict, search_item: dict):
    return date_it(first_non_empty(detail_body.get("dataPubblicazione"), detail_body.get("dtPubblicazione"), search_item.get("dataPubblicazione"), search_item.get("dtPubblicazione")))


def best_descrizione(search_item: dict, detail_body: dict) -> str:
    for v in (search_item.get("descLotto"), (detail_body.get("lotto") or {}).get("descLotto"), detail_body.get("descLotto")):
        s = clean(v)
        if s:
            return s
    return ""


def structured_field_from_beni(detail_body: dict, keys, suffix=""):
    vals = []
    beni = (detail_body.get("beni") or []) + ((detail_body.get("lotto") or {}).get("beni") or [])
    for b in beni:
        for k in keys:
            s = clean(b.get(k))
            if s and s not in vals:
                vals.append(s)
    if not vals:
        return ""
    joined = "+".join(vals)
    return f"{joined} {suffix}".strip()


def extract_superficie_text(text: str, descr: str) -> str:
    blob = " ".join([text or "", descr or ""])
    vals = []
    for s in re.findall(r"\b(\d{1,6}(?:[.,]\d{1,2})?)\s*(?:mq|m2|m²)\b", blob, flags=re.I):
        n = s.replace(".", ",")
        if n not in vals:
            vals.append(n)
    return f"{'+'.join(vals)} mq" if vals else ""


def extract_vani_text(text: str, descr: str) -> str:
    blob = " ".join([text or "", descr or ""])
    m = re.search(r"\b(\d{1,2})(?:[.,](\d))?\s*vani\b", blob, flags=re.I)
    return m.group(1) + (("," + m.group(2)) if m and m.group(2) else "") if m else ""


def extract_piano_text(text: str, descr: str) -> str:
    blob = " ".join([text or "", descr or ""]).lower()
    if re.search(r"\bpiano\s+terra\b|\bpt\b|\bp\.t\.\b", blob):
        return "Piano terra"
    if "rialzato" in blob and "piano" in blob:
        return "Piano rialzato"
    if "seminterr" in blob:
        return "Seminterrato"
    m = re.search(r"\b(\d{1,2})\s*(?:°|º)?\s*piano\b", blob)
    return m.group(1) if m else ""


def note_from_raw(raw_text: str, descrizione: str) -> str:
    text = " ".join([raw_text or "", descrizione or ""])
    kw = re.compile(r"(abus|difform|irregolar|sanator|condon|titolo abilitativo|assenza di titolo|privo di titolo)", re.I)
    chunks = re.split(r"(?<=[\.!?;])\s+|\n+", text)
    out = []
    for ch in chunks:
        s = clean(ch)
        if s and kw.search(s) and s not in out:
            out.append(s)
    return " | ".join(out)


def derive_listing_row(obj, allegati_by_id):
    pvp_id = clean(obj.get("pvp_id"))
    link_pvp = clean(obj.get("link_pvp"))
    raw_text = obj.get("raw_text") or ""
    search_item = obj.get("search_item") or {}
    detail_body = obj.get("detail_body") or {}
    descrizione = best_descrizione(search_item, detail_body)
    bene_list = (detail_body.get("beni") or []) + ((detail_body.get("lotto") or {}).get("beni") or [])
    geo_api = extract_geo_api(detail_body, search_item, bene_list)
    indirizzo = ""
    for bene in bene_list:
        indirizzo = build_address_from_bene(bene)
        if indirizzo:
            break
    if not indirizzo:
        s_addr = search_item.get("indirizzo") or {}
        comune0 = clean(s_addr.get("descComune") or s_addr.get("citta") or s_addr.get("comune"))
        provincia0 = clean(s_addr.get("descProvincia") or s_addr.get("provincia"))
        sigla0 = clean(s_addr.get("siglaProvincia") or s_addr.get("sigla_provincia"))
        cap0 = re.sub(r"\D", "", clean(s_addr.get("cap") or "")).zfill(5) if clean(s_addr.get("cap") or "") else ""
        via0 = clean_via_structured(s_addr.get("via") or s_addr.get("descVia"), comune0, provincia0, sigla0, cap0)
        nc0 = clean(s_addr.get("numeroCivico"))
        if nc0 and nc0 not in via0:
            via0 = " ".join(x for x in [via0, nc0] if x).strip()
        indirizzo = compose_indirizzo(via0, cap0, comune0, sigla0 or provincia0)
    info = match_comune_info(indirizzo)
    comune = first_non_empty(geo_api.get("comune"), info.get("comune"))
    provincia = first_non_empty(geo_api.get("provincia"), info.get("provincia"))
    sigla = first_non_empty(geo_api.get("sigla_provincia"), info.get("sigla_provincia"))
    regione = first_non_empty(geo_api.get("regione"), info.get("regione"))
    cap = first_non_empty(geo_api.get("cap"), info.get("cap"))
    geo_norm = normalize_geo_with_csv(comune, provincia, sigla, regione, cap)
    comune = geo_norm.get("comune", "")
    provincia = geo_norm.get("provincia", "")
    sigla = geo_norm.get("sigla_provincia", "")
    regione = geo_norm.get("regione", "")
    cap = geo_norm.get("cap", "")

    geo_canon = canonicalize_geo_labels(regione, provincia, sigla)
    provincia = geo_canon.get("provincia", "")
    sigla = geo_canon.get("sigla_provincia", "")
    regione = geo_canon.get("regione", "")
    if bene_list:
        addr = (bene_list[0].get("indirizzo") or {})
        via_final = clean_via_structured(addr.get("via") or addr.get("descVia"), comune, provincia, sigla, cap)
        nc_final = clean(addr.get("numeroCivico"))
        if nc_final and nc_final not in via_final:
            via_final = " ".join(x for x in [via_final, nc_final] if x).strip()
        indirizzo = compose_indirizzo(via_final, cap, comune, sigla or provincia)
    else:
        s_addr = search_item.get("indirizzo") or {}
        via_final = clean_via_structured(s_addr.get("via") or s_addr.get("descVia"), comune, provincia, sigla, cap)
        nc_final = clean(s_addr.get("numeroCivico"))
        if nc_final and nc_final not in via_final:
            via_final = " ".join(x for x in [via_final, nc_final] if x).strip()
        if via_final or comune or cap:
            indirizzo = compose_indirizzo(via_final, cap, comune, sigla or provincia)
    tipologia = pick_tipologia(detail_body, search_item, "")
    categoria_immobile, categoria_immobile_reason = pick_categoria_immobile(detail_body, search_item, "")
    categoria_immobile = categoria_immobile or "altro"
    procedura = detail_body.get("procedura") or {}
    tipo_procedura = pick_tipo_procedura(detail_body, search_item, "")
    categoria_procedura, categoria_procedura_reason = pick_categoria_procedura(detail_body, search_item, "")
    numero_procedura = re.sub(r"\D+", "", clean(first_non_empty(procedura.get("numeRg"), detail_body.get("numeRg"), search_item.get("numeRg"))))
    anno_procedura = re.sub(r"\D+", "", clean(first_non_empty(procedura.get("numeAnnoRg"), detail_body.get("numeAnnoRg"), search_item.get("numeAnnoRg"))))[:4]
    prezzo_base = euro_it(first_non_empty(search_item.get("prezzoBaseAsta"), detail_body.get("impoBaseAsta")))
    offerta_minima = euro_it(first_non_empty(search_item.get("offertaMinima"), detail_body.get("impoOffertaMinima")))
    rialzo_minimo = euro_it(first_non_empty(search_item.get("rialzoMinimo"), detail_body.get("impoOffertaAumento")))
    data_asta = date_it(first_non_empty(search_item.get("dataVendita"), detail_body.get("dataVendita")))
    data_pubblicazione = pick_data_pubblicazione(detail_body, search_item)
    disponibilita = normalize_disponibilita(first_non_empty(search_item.get("disponibilita"), detail_body.get("disponibilita")))
    if not disponibilita:
        for bene in bene_list:
            disponibilita = normalize_disponibilita(first_non_empty(bene.get("disponibilita"), bene.get("disponibilitaDesc")))
            if disponibilita:
                break
    tribunale = first_non_empty(search_item.get("tribunale"), procedura.get("descUfficio"), detail_body.get("tribunale"))
    lotto = first_non_empty(search_item.get("numeroLotto"), (detail_body.get("lotto") or {}).get("numeroLotto"), (detail_body.get("lotto") or {}).get("codiLotto"))
    superficie_mq = structured_field_from_beni(detail_body, ["superficie", "superficieCommerciale", "superficieMq"], "mq") or extract_superficie_text(raw_text, descrizione)
    vani = structured_field_from_beni(detail_body, ["numeroVani", "vani"]) or extract_vani_text(raw_text, descrizione)
    piano = structured_field_from_beni(detail_body, ["piano", "descPiano", "ubicazionePiano"]) or extract_piano_text(raw_text, descrizione)
    filenames = [clean(a.get("filename")) for a in allegati_by_id.get(pvp_id, []) if clean(a.get("filename"))]
    esperimento = clean_esperimento(first_non_empty(detail_body.get("numeEsperimento"), search_item.get("numeEsperimento")) or extract_esperimento_from_filenames(filenames))
    lat = clean(geo_api.get("lat"))
    lon = clean(geo_api.get("lon"))
    maps_link = build_maps_link(lat, lon, indirizzo, comune, provincia)
    return {
        "pvp_id": pvp_id,
        "link_pvp": link_pvp,
        "data_pubblicazione": data_pubblicazione,
        "data_pubblicazione_date": iso_date_from_it(data_pubblicazione),
        "data_asta": data_asta,
        "regione": regione,
        "comune": comune,
        "provincia": provincia,
        "sigla_provincia": sigla,
        "cap": cap,
        "indirizzo": indirizzo,
        "maps_link": maps_link,
        "latitudine": lat,
        "longitudine": lon,
        "tipologia": tipologia,
        "tribunale": tribunale,
        "tipo_procedura": tipo_procedura,
        "numero_procedura": numero_procedura,
        "anno_procedura": anno_procedura,
        "categoria_procedura": categoria_procedura,
        "categoria_procedura_reason": categoria_procedura_reason,
        "lotto": lotto,
        "descrizione": descrizione,
        "categoria_immobile": categoria_immobile,
        "categoria_immobile_reason": categoria_immobile_reason,
        "prezzo_base": prezzo_base,
        "offerta_minima": offerta_minima,
        "rialzo_minimo": rialzo_minimo,
        "disponibilita": disponibilita,
        "vani": vani,
        "piano": piano,
        "superficie_mq": superficie_mq,
        "esperimento_vendita": esperimento,
        "note": note_from_raw(raw_text, descrizione),
        "status": "active",
    }


def _to_int(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v)
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        ds = re.sub(r"\D", "", s)
        return int(ds) if ds else None
    return None


def _http_json(method: str, url: str, payload=None, prefer=None):
    data = None
    headers = {
        "apikey": SUPABASE_SERVICE_ROLE,
        "Authorization": f"Bearer {SUPABASE_SERVICE_ROLE}",
    }
    if payload is not None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"
    if prefer:
        headers["Prefer"] = prefer
    req = Request(url, data=data, method=method)
    for k, v in headers.items():
        req.add_header(k, v)
    with urlopen(req, timeout=120) as resp:
        return resp.read().decode("utf-8", errors="replace")


def sanitize_listing_rows(rows):
    BIGINT_KEYS = {"pvp_id", "numero_procedura", "anno_procedura"}
    def _looks_int_key(k: str) -> bool:
        k2 = (k or "").lower()
        return k2 in BIGINT_KEYS or k2.endswith("_id") or "numero" in k2 or "anno" in k2
    out = []
    for r in rows:
        rr = {k: r.get(k) for k in LISTINGS_FIELDS if k in r}
        for k in list(rr.keys()):
            v = rr.get(k)
            if isinstance(v, str) and v.strip() == "":
                rr.pop(k, None)
        for k in list(rr.keys()):
            v = rr.get(k)
            if k in {"pvp_id", "numero_procedura", "anno_procedura"} or (isinstance(v, str) and _looks_int_key(k)):
                iv = _to_int(v)
                if iv is None:
                    rr.pop(k, None)
                else:
                    rr[k] = iv
        if "pvp_id" not in rr:
            continue
        out.append(rr)
    return out


def chunked(seq, n):
    for i in range(0, len(seq), n):
        yield seq[i:i+n]




def upsert_allegati(rows):
    if not SUPABASE_ENABLED or not rows:
        return
    clean_rows = []
    for r in rows:
        rr = {
            "pvp_id": clean(r.get("pvp_id")),
            "filename": clean(r.get("filename")),
            "resource_url": clean(r.get("resource_url")),
            "mime": clean(r.get("mime")),
            "size_bytes": _to_int(r.get("size_bytes")),
            "seen_at": clean(r.get("seen_at")),
        }
        if not rr["pvp_id"] or not rr["filename"] or not rr["resource_url"]:
            continue
        for k in list(rr.keys()):
            if rr[k] in ("", None):
                rr.pop(k, None)
        clean_rows.append(rr)
    if not clean_rows:
        return
    url = f"{SUPABASE_URL.rstrip('/')}/rest/v1/listing_allegati?on_conflict=pvp_id,resource_url"
    from collections import defaultdict as dd
    groups = dd(list)
    for r in clean_rows:
        groups[tuple(sorted(r.keys()))].append(r)
    for _, batch_rows in groups.items():
        for batch in chunked(batch_rows, BATCH_SIZE_ALLEGATI):
            _http_json("POST", url, batch, prefer="resolution=merge-duplicates,return=minimal")




def insert_soggetti(rows):
    if not SUPABASE_ENABLED or not rows:
        return
    clean_rows = []
    for r in rows:
        rr = {
            "pvp_id": clean(r.get("pvp_id")),
            "role_code": clean(r.get("role_code")),
            "role_label": clean(r.get("role_label")),
            "nome": clean(r.get("nome")),
            "cognome": clean(r.get("cognome")),
            "display_name": clean(r.get("display_name")),
            "email": clean(r.get("email")),
            "telefono": clean(r.get("telefono")),
            "sort_order": _to_int(r.get("sort_order")),
        }
        if not rr["pvp_id"]:
            continue
        for k in list(rr.keys()):
            if rr[k] in ("", None):
                rr.pop(k, None)
        clean_rows.append(rr)
    from collections import defaultdict as dd
    groups = dd(list)
    for r in clean_rows:
        groups[tuple(sorted(r.keys()))].append(r)
    url = f"{SUPABASE_URL.rstrip('/')}/rest/v1/listing_soggetti"
    for _, batch_rows in groups.items():
        for batch in chunked(batch_rows, BATCH_SIZE_ALLEGATI):
            _http_json("POST", url, batch, prefer="return=minimal")


def normalize_soggetti_rows(raw_soggetti_rows):
    out = []
    for r in raw_soggetti_rows:
        pvp_id = clean(r.get("pvp_id"))
        if not pvp_id:
            continue
        out.append({
            "pvp_id": pvp_id,
            "role_code": clean(r.get("role_code") or r.get("ruolo_code") or r.get("codiRuolo")),
            "role_label": clean(r.get("role_label") or r.get("ruolo_label") or r.get("descRuolo")),
            "nome": clean(r.get("nome")),
            "cognome": clean(r.get("cognome")),
            "display_name": clean(r.get("display_name") or r.get("nominativo") or " ".join(x for x in [clean(r.get("nome")), clean(r.get("cognome"))] if x)),
            "email": clean(r.get("email")),
            "telefono": clean(r.get("telefono")),
            "sort_order": r.get("sort_order") if r.get("sort_order") is not None else r.get("ordine"),
        })
    seen = set()
    dedup = []
    for r in out:
        key = (r["pvp_id"], r.get("role_code", ""), r.get("display_name", ""), r.get("email", ""))
        if key in seen:
            continue
        seen.add(key)
        dedup.append(r)
    return dedup






def _norm_admin_text(v):
    s = clean(v)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.casefold().replace("’", "'")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


CURRENT_COMUNE_ALIASES = {
    "presicce": "presicce acquarica",
    "molvena": "colceresa",
    "vermezzo": "vermezzo con zelo",
    "corigliano calabro": "corigliano rossano",
    "rossano": "corigliano rossano",
    "popoli": "popoli terme",
    "lusiana": "lusiana conco",
    "campospinoso": "campospinoso albaredo",
    "veruno": "gattico veruno",
    "tavarnelle val di pesa": "barberino tavarnelle",
    "montebello ionico": "montebello jonico",
    "tonengo": "moransengo tonengo",
    "vighizzolo d este": "borgo veneto",
    "gambugliano": "sovizzo",
    "piadena": "piadena drizzona",
}


def load_istat_admin(path):
    """
    Lookup amministrativo corrente da Elenco-comuni-italiani.xlsx.
    Fail-closed: se non si riesce a leggere il file, non inventa sigle.
    """
    if not path.exists():
        return {"ok": False, "reason": "istat_xlsx_missing", "by_comune": {}}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        header = None
        header_row = None
        for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            norms = [_norm_admin_text(v) for v in row]
            if any("sigla automobilistica" in x for x in norms) and any("denominazione" in x for x in norms):
                header = norms
                header_row = ridx
                break
            if ridx >= 15:
                break
        if not header:
            return {"ok": False, "reason": "istat_header_not_found", "by_comune": {}}

        def first_idx(tests):
            for test in tests:
                for i, h in enumerate(header):
                    if test(h):
                        return i
            return None

        idx_comune = first_idx([
            lambda h: h == "denominazione in italiano",
            lambda h: "denominazione italiana e straniera" in h,
            lambda h: h.startswith("denominazione") and "unita territoriale" not in h and "regione" not in h,
        ])
        idx_sigla = first_idx([lambda h: "sigla automobilistica" in h])
        idx_prov = first_idx([
            lambda h: "denominazione dell unita territoriale sovracomunale" in h,
            lambda h: "denominazione provincia" in h,
            lambda h: h == "provincia",
        ])
        idx_reg = first_idx([
            lambda h: h == "regione",
            lambda h: "denominazione regione" in h,
        ])
        if idx_comune is None or idx_sigla is None or idx_prov is None or idx_reg is None:
            return {
                "ok": False,
                "reason": f"istat_columns_missing comune={idx_comune} sigla={idx_sigla} prov={idx_prov} reg={idx_reg}",
                "by_comune": {},
            }

        by_comune = {}
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            def at(idx):
                return clean(row[idx]) if idx is not None and idx < len(row) else ""
            comune = at(idx_comune)
            sigla = at(idx_sigla).upper()
            provincia = at(idx_prov)
            regione = at(idx_reg)
            if not comune:
                continue
            rec = {
                "comune": comune,
                "sigla_provincia": sigla,
                "provincia": provincia,
                "regione": regione,
            }
            by_comune.setdefault(_norm_admin_text(comune), []).append(rec)
        return {"ok": True, "reason": "ok", "by_comune": by_comune}
    except Exception as e:
        return {"ok": False, "reason": f"istat_load_error:{type(e).__name__}", "by_comune": {}}


def canonical_admin(row, istat):
    comune_raw = clean(row.get("comune"))
    provincia_raw = clean(row.get("provincia"))
    sigla_raw = clean(row.get("sigla_provincia")).upper()
    regione_raw = clean(row.get("regione"))
    key = _norm_admin_text(comune_raw)
    key = CURRENT_COMUNE_ALIASES.get(key, key)

    cands = list((istat.get("by_comune") or {}).get(key, []))
    if not cands:
        return {
            "ok": False,
            "reason": "istat_comune_not_found",
            "comune": comune_raw,
            "provincia": provincia_raw,
            "sigla_provincia": sigla_raw,
            "regione": regione_raw,
        }

    if len(cands) > 1:
        pkey = _norm_admin_text(provincia_raw)
        skey = sigla_raw
        narrowed = [
            r for r in cands
            if (pkey and _norm_admin_text(r.get("provincia")) == pkey)
            or (skey and clean(r.get("sigla_provincia")).upper() == skey)
        ]
        if len(narrowed) == 1:
            cands = narrowed

    if len(cands) != 1:
        return {
            "ok": False,
            "reason": f"istat_comune_ambiguous:{len(cands)}",
            "comune": comune_raw,
            "provincia": provincia_raw,
            "sigla_provincia": sigla_raw,
            "regione": regione_raw,
        }

    r = cands[0]
    return {
        "ok": bool(clean(r.get("sigla_provincia"))),
        "reason": "istat_resolved" if clean(r.get("sigla_provincia")) else "istat_sigla_missing",
        "comune": clean(r.get("comune")) or comune_raw,
        "provincia": clean(r.get("provincia")) or provincia_raw,
        "sigla_provincia": clean(r.get("sigla_provincia")).upper(),
        "regione": clean(r.get("regione")) or regione_raw,
    }


def parse_auction_date(v):
    s = clean(v)
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    return None


def classify_new_candidate(row, today, istat):
    """
    Lifecycle Aste Easy:
    - data asta oggi/futura: inseribile;
    - scaduta da 1..14 giorni: ancora inseribile, per la finestra di retention;
    - scaduta da oltre 14 giorni: non reinserire;
    - data asta non leggibile: review manuale, mai insert automatico.
    """
    d = parse_auction_date(row.get("data_asta"))
    if d is None:
        return "SKIP_DATE_UNKNOWN", "auction_date_missing_or_unparseable", None

    age_days = (today - d).days

    future_days = -age_days
    if future_days > MAX_FUTURE_AUCTION_DAYS:
        return (
            "SKIP_DATE_FAR_FUTURE_REVIEW",
            f"auction_{future_days}_days_in_future_gt_{MAX_FUTURE_AUCTION_DAYS}",
            None,
        )

    if age_days > RETENTION_DAYS:
        return (
            "SKIP_EXPIRED_BEYOND_RETENTION",
            f"auction_expired_{age_days}_days_gt_{RETENTION_DAYS}",
            None,
        )

    admin = canonical_admin(row, istat)
    if not admin.get("ok"):
        return "SKIP_ADMIN_REVIEW", clean(admin.get("reason")), admin

    if age_days > 0:
        return (
            "ELIGIBLE_EXPIRED_WITHIN_RETENTION",
            f"auction_expired_{age_days}_days_within_{RETENTION_DAYS}",
            admin,
        )

    return "ELIGIBLE_ACTIVE", "auction_today_or_future_and_admin_resolved", admin




# ----------------- Importazione ordinaria consolidata: SOLO NUOVI ID -----------------
# Nessuna dipendenza da pvp_import_incremental_v12_3.py. Nessun geocoding/API Google.
FETCH_BATCH = 150
RETENTION_DAYS = 14
MAX_FUTURE_AUCTION_DAYS = 730
ISTAT_XLSX = REFERENCE_DIR / 'Elenco-comuni-italiani.xlsx'


def coordinate_pair(addr):
    """Only accept a complete pair from the *same* PVP address object."""
    c = addr.get('coordinate') or {}
    if not isinstance(c, dict): c = {}
    lat = first_non_empty(addr.get('latitudine'), addr.get('lat'), c.get('latitudine'), c.get('lat'))
    lon = first_non_empty(addr.get('longitudine'), addr.get('lon'), addr.get('lng'), c.get('longitudine'), c.get('lon'), c.get('lng'))
    try:
        x,y = float(str(lat).replace(',', '.')), float(str(lon).replace(',', '.'))
        if math.isfinite(x) and math.isfinite(y) and -90 <= x <= 90 and -180 <= y <= 180 and not (x == 0 and y == 0):
            return x,y
    except (TypeError, ValueError):
        pass
    return None,None


def primary_geo(obj):
    """Never splice geography across separate beni/lotto/search addresses."""
    detail=obj.get('detail_body') or {}
    search=obj.get('search_item') or {}
    lot=detail.get('lotto') or {}
    beni=list(detail.get('beni') or [])+list(lot.get('beni') or [])
    candidates=[]
    for b in beni:
        if not isinstance(b, dict): continue
        addr=b.get('indirizzo') or {}
        if isinstance(addr,dict): candidates.append(('bene',addr))
    for kind,container in (('search',search),('detail',detail)):
        addr=container.get('indirizzo') or {}
        if isinstance(addr,dict): candidates.append((kind,addr))
    selected=None
    # Prefer the first bene with a named municipality; otherwise search/detail.
    for kind,addr in candidates:
        if first_non_empty(addr.get('descComune'),addr.get('citta'),addr.get('comune')):
            selected=(kind,addr);break
    if selected is None:
        return {'source':'unresolved','comune':'','provincia':'','sigla_provincia':'',
                'regione':'','cap':'','indirizzo':'','latitudine':None,'longitudine':None}
    kind,addr=selected
    comune=first_non_empty(addr.get('descComune'),addr.get('citta'),addr.get('comune'))
    provincia=first_non_empty(addr.get('descProvincia'),addr.get('provincia'))
    sigla=first_non_empty(addr.get('siglaProvincia'),addr.get('sigla_provincia')).upper()
    regione=first_non_empty(addr.get('descRegione'),addr.get('regione'),addr.get('codRegione'))
    cap=re.sub(r'\D','',clean(addr.get('cap'))).zfill(5) if clean(addr.get('cap')) else ''
    via=clean_via_structured(addr.get('via') or addr.get('descVia'), comune, provincia, sigla, cap)
    civ=clean(addr.get('numeroCivico'))
    if civ and civ not in via: via=(via+' '+civ).strip()
    lat,lon=coordinate_pair(addr)
    return {'source':kind,'comune':comune,'provincia':provincia,'sigla_provincia':sigla,
            'regione':regione,'cap':cap,'indirizzo':compose_indirizzo(via,cap,comune,sigla or provincia),
            'latitudine':lat,'longitudine':lon}


def derive_new_consistent(obj, allegati_by_id):
    row=derive_listing_row(obj,allegati_by_id)
    geo=primary_geo(obj)
    for key in ('comune','provincia','sigla_provincia','regione','cap','indirizzo','latitudine','longitudine'):
        row[key]=geo[key]
    row['maps_link']=build_maps_link(geo['latitudine'],geo['longitudine'],geo['indirizzo'],geo['comune'],geo['provincia'])
    return row,geo['source']


def fetch_existing_ids(base,key,ids):
    """PVP ID-only GET: avoids reading or deriving existing listings' full content."""
    if any(not re.fullmatch(r'[0-9]{1,19}',str(pid)) for pid in ids):
        raise ValueError('pvp_id input non valido')
    out=set()
    for off in range(0,len(ids),FETCH_BATCH):
        group=ids[off:off+FETCH_BATCH]
        if not group: continue
        url=base.rstrip('/')+'/rest/v1/listings?'+urlencode({'select':'pvp_id','pvp_id':'in.('+','.join(group)+')','limit':str(FETCH_BATCH)})
        req=Request(url,method='GET',headers={'apikey':key,'Authorization':'Bearer '+key,'Accept':'application/json'})
        with urlopen(req,timeout=120) as resp:
            data=json.loads(resp.read().decode('utf-8') or '[]')
        if not isinstance(data,list) or len(data)>len(group):
            raise RuntimeError('Supabase: lista esistenti inattesa')
        got=[str(r.get('pvp_id')) for r in data]
        if len(got)!=len(set(got)) or not set(got).issubset(set(group)):
            raise RuntimeError('Supabase: ID esistenti duplicati/inattesi')
        out.update(got)
    return out


def plan_new_only(raw_rows,allegati_rows,soggetti_rows,existing_ids,istat,today):
    raw_by_id={}
    duplicates=[]
    for obj in raw_rows:
        pid=clean(obj.get('pvp_id'))
        if not re.fullmatch(r'[0-9]{1,19}',pid):
            raise ValueError('RAW: pvp_id non valido')
        if pid in raw_by_id:
            duplicates.append(pid)
            # Fail-closed: different payload for same pvp_id is not silently overwritten.
            if obj!=raw_by_id[pid]:
                raise ValueError('RAW: dati diversi per lo stesso pvp_id '+pid)
        raw_by_id[pid]=obj
    allegati_by_id=defaultdict(list)
    for a in allegati_rows:
        allegati_by_id[clean(a.get('pvp_id'))].append(a)
    eligible=[]; decisions=Counter(); sources=Counter(); review=[]
    for pid,obj in raw_by_id.items():
        if pid in existing_ids: continue
        row,source=derive_new_consistent(obj,allegati_by_id)
        decision,reason,admin=classify_new_candidate(row,today,istat)
        decisions[decision]+=1
        if decision.startswith('ELIGIBLE_'):
            row.update({k:admin[k] for k in ('comune','provincia','sigla_provincia','regione')})
            row['maps_link']=build_maps_link(row['latitudine'],row['longitudine'],row['indirizzo'],row['comune'],row['provincia'])
            eligible.append(row);sources[source]+=1
        else:
            review.append({'pvp_id':pid,'decision':decision,'reason':reason})
    eligible_set={clean(r['pvp_id']) for r in eligible}
    subjects=[r for r in normalize_soggetti_rows(soggetti_rows) if clean(r['pvp_id']) in eligible_set]
    attachments=[r for r in allegati_rows if clean(r.get('pvp_id')) in eligible_set]
    return {'new_candidates':len(raw_by_id)-len(set(raw_by_id)&existing_ids),
            'eligible':eligible,'decisions':decisions,'sources':sources,'review':review,
            'subjects':subjects,'attachments':attachments,'duplicates':len(duplicates)}


def consolidated_main(argv=None):
    ap=argparse.ArgumentParser(description='Aste Easy - unico importer, solo nuovi ID; anteprima senza scritture')
    ap.add_argument('--apply',action='store_true',help='non ancora abilitato: consolidamento geocoder/stato DB in corso')
    args=ap.parse_args(argv)
    if args.apply:
        ap.error('APPLY BLOCCATO in questa anteprima: usare la pipeline attuale finche non sara pronto il geocoder unico')
    ensure_runtime_dirs()
    for path in (IN_RAW,IN_ALLEGATI,IN_SOGGETTI,ISTAT_XLSX):
        if not path.is_file(): raise SystemExit(f'Input necessario mancante: {path}')
    load_project_env(required=('SUPABASE_URL','SUPABASE_SERVICE_ROLE'))
    base=os.environ['SUPABASE_URL'].strip();key=os.environ['SUPABASE_SERVICE_ROLE'].strip()
    raw=load_jsonl(IN_RAW)
    allegati=load_jsonl(IN_ALLEGATI)
    soggetti=load_jsonl(IN_SOGGETTI)
    ids=sorted({clean(r.get('pvp_id')) for r in raw})
    existing=fetch_existing_ids(base,key,ids)
    istat=load_istat_admin(ISTAT_XLSX)
    if not istat.get('ok'):
        raise SystemExit('ISTAT non disponibile: '+str(istat.get('reason')))
    plan=plan_new_only(raw,allegati,soggetti,existing,istat,datetime.now().astimezone().date())
    # One compact summary; no per-ID manifest needed by the planned DB-driven geocoder.
    summary={'mode':'GET_ONLY_IMPORT_PREVIEW_ZERO_GOOGLE_ZERO_PATCH','incoming_unique':len(ids),
        'existing':len(existing),'new_candidates':plan['new_candidates'],
        'eligible_new':len(plan['eligible']),'new_soggetti':len(plan['subjects']),
        'new_allegati':len(plan['attachments']),'decision_counts':dict(plan['decisions']),
        'primary_address_sources':dict(plan['sources']),
        'eligible_with_pvp_coords':sum(r['latitudine'] is not None and r['longitudine'] is not None for r in plan['eligible']),
        'eligible_without_pvp_coords':sum(r['latitudine'] is None or r['longitudine'] is None for r in plan['eligible']),
        'duplicate_input_ids_same_payload':plan['duplicates'],
        'existing_writes':0,'db_patches':0,'google_calls':0,
        'raw_sha256':hashlib.sha256(IN_RAW.read_bytes()).hexdigest(),
        'notes':'Anteprima architetturale: non sostituisce ancora la pipeline operativa; geocoder/stato DB da consolidare.'}
    stamp=datetime.now().astimezone().strftime('%Y%m%d-%H%M%S')
    dest=REPORTS_DIR / f'pvp_single_importer_preview_v1_{stamp}.json'
    dest.write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    print('[OUT]',dest)
    return 0

if __name__ == '__main__':
    raise SystemExit(consolidated_main())
